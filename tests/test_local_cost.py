"""Days and Local cost in Matched / Only in MR / Only in CTR. Days (MR) / Days (CTR): the days of that side's line (equipment
only). Local cost: SAGE's local_expect_cost; with the toggle, the TOTAL by the statistics' formula — equipment = days x cost x
quantity, consumables = cost x quantity. Fixed fixture costs (the test SAGE export has random ones): the finished statistics are
injected into the panel."""
import pandas as pd
from PySide6.QtCore import Qt

import harness as H
import test_mismatch_navigation as T
from ctr_tools import stockcode_types as ST
from stats import cost_stats as cs

COSTS = {"C1": 10.0, "C2": 20.0, "C3": 30.0}          # C4, C5 are not in "SAGE"


def _names(view):
    m = view.model()
    return [m.headerData(c, Qt.Orientation.Horizontal) for c in range(m.columnCount())]


def _col(view, name):
    m = view.model()
    c = _names(view).index(name)
    return [m.index(r, c).data() for r in range(m.rowCount())]


def _tip(view, row, name):
    m = view.model()
    c = _names(view).index(name)
    return m.index(row, c).data(Qt.ItemDataRole.ToolTipRole)


def _window_with_costs(days=None):
    # (the fixture CTR says "Plant & Equipment" for every code, so the types are confirmed here)
    ST.save_types({"equipment": {"C1", "C2"}, "consumable": {"C3", "C4", "C5"},
                   "confirmed": {"C1": "Equipment", "C2": "Equipment", "C3": "Consumable",
                                 "C4": "Consumable", "C5": "Consumable"}})
    win = T._window()                                   # 5 matched codes (C2 and C4 twice) + the stats job has finished
    win._stats_panel.set_rows(cs.cost_rows(win._stats_items, COSTS, True))
    if days:
        win._stats_panel.set_days(days)
    win._local_ready = True                             # what the app does when the statistics have finished
    win._apply_local_costs()
    H.pump(5)
    return win


def _ctx(unit, info):
    return {"unit": unit, "info": info}


# ── the pure function ────────────────────────────────────────────────────────

def test_the_unit_column_and_the_days_columns_end_the_table():
    import app as A
    df = pd.DataFrame({"Stock Code": ["c1", "C4"], "Qty (MR)": [2, 3], "Qty (CTR)": [5, None],
                       "MR Document": ["mr", "mr"], "MR Sheet": ["S", "S"], "CTR Document": ["c", ""], "CTR Sheet": ["P", ""]})
    ctx = _ctx(COSTS, {("MR", "mr", "S", "C1"): (cs.EQUIPMENT, 4.0, ""), ("CTR", "c", "P", "C1"): (cs.EQUIPMENT, 4.0, "")})
    out, _tips = A.MainWindow._with_local_costs(df, ctx, total=False)
    assert list(out.columns)[-3:] == ["Days (MR)", "Days (CTR)", "Local cost"]
    assert out["Local cost"].tolist()[0] == 10.0 and pd.isna(out["Local cost"].tolist()[1])          # C4 is not in SAGE
    assert out["Days (MR)"].tolist()[0] == 4.0 and pd.isna(out["Days (MR)"].tolist()[1])


def test_the_total_is_days_x_cost_x_qty_for_equipment_and_cost_x_qty_for_consumables():
    import app as A
    df = pd.DataFrame({"Stock Code": ["C1", "C3"], "Qty (MR)": [2, 3], "Qty (CTR)": [5, 4],
                       "MR Document": ["mr", "mr"], "MR Sheet": ["S", "S"], "CTR Document": ["c", "c"], "CTR Sheet": ["P", "P"]})
    ctx = _ctx(COSTS, {("MR", "mr", "S", "C1"): (cs.EQUIPMENT, 3.0, ""), ("CTR", "c", "P", "C1"): (cs.EQUIPMENT, 6.0, ""),
                       ("MR", "mr", "S", "C3"): (cs.CONSUMABLE, float("nan"), ""), ("CTR", "c", "P", "C3"): (cs.CONSUMABLE, float("nan"), "")})
    out, _ = A.MainWindow._with_local_costs(df, ctx, total=True)
    assert list(out.columns)[-4:] == ["Days (MR)", "Days (CTR)", "Local cost total (MR)", "Local cost total (CTR)"]
    assert out["Local cost total (MR)"].tolist() == [3 * 10 * 2, 30 * 3]                 # equipment: days x cost x qty; consumable: cost x qty
    assert out["Local cost total (CTR)"].tolist() == [6 * 10 * 5, 30 * 4]
    assert pd.isna(out["Days (MR)"].tolist()[1])                                          # a consumable shows no days
    again, _ = A.MainWindow._with_local_costs(out, ctx, total=False)                     # earlier columns are replaced, not stacked
    assert [c for c in again.columns if str(c).startswith(("Local cost", "Days ("))] == ["Days (MR)", "Days (CTR)", "Local cost"]


def test_a_total_that_cannot_be_computed_stays_blank_and_says_why():
    import app as A
    keys = ["C1", "C2", "C3", "C4", "X9"]
    df = pd.DataFrame({"Stock Code": keys, "Qty (MR)": [1, 1, None, 1, 1], "Qty (CTR)": [None] * 5,
                       "MR Document": "mr", "MR Sheet": "S", "CTR Document": "", "CTR Sheet": ""})
    ctx = _ctx({"C1": 10.0, "C2": 20.0, "C3": 30.0, "X9": 5.0},
               {("MR", "mr", "S", "C1"): (cs.EQUIPMENT, float("nan"), ""),            # equipment, no days
                ("MR", "mr", "S", "C2"): (cs.CONSUMABLE, float("nan"), "conflict"),   # type undecided
                ("MR", "mr", "S", "C3"): (cs.CONSUMABLE, float("nan"), ""),           # no quantity
                ("MR", "mr", "S", "C4"): (cs.CONSUMABLE, float("nan"), ""),           # not in SAGE
                ("MR", "mr", "S", "X9"): ("", float("nan"), "unknown")})              # type unknown
    out, tips = A.MainWindow._with_local_costs(df, ctx, total=True)
    assert out["Local cost total (MR)"].isna().all()
    why = {keys[r]: t for (r, c), t in tips.items() if c == "Local cost total (MR)"}
    assert "days are missing" in why["C1"] and "decision" in why["C2"] and "quantity" in why["C3"]
    assert "SAGE" in why["C4"] and "unknown" in why["X9"]
    assert "Days" in tips[(0, "Days (MR)")]                                              # the Days cell says it too


# ── the tables ───────────────────────────────────────────────────────────────

def test_days_and_local_cost_are_the_last_columns_of_all_three_tables_and_blank_until_the_statistics_exist():
    win = T._window()
    for view in (win._tab_matched, win._tab_only_mr, win._tab_only_ctr):
        if view.model() is None:
            continue
        assert _names(view)[-3:] == ["Days (MR)", "Days (CTR)", "Local cost"], _names(view)
        assert set(_col(view, "Local cost")) <= {"", "⏳ waiting"}               # shown as waiting, never as a number
    assert win._display_df["Local cost"].isna().all()                         # and the data underneath is blank


def test_the_tables_show_days_and_sages_cost_in_matched():
    win = _window_with_costs(days={"C1": 3.0})
    names = _names(win._tab_matched)
    assert names[-1] == "Local cost" and names.index("Days (MR)") > names.index("CTR Sheet")
    rows = dict(zip(_col(win._tab_matched, "Stock Code"), zip(_col(win._tab_matched, "Days (MR)"),
                                                              _col(win._tab_matched, "Days (CTR)"),
                                                              _col(win._tab_matched, "Local cost"))))
    assert rows["C1"] == ("3.0", "3.0", "10.0")                   # equipment: the days typed in Statistics, on both sides
    assert rows["C3"] == ("", "", "30.0")                         # a consumable has no days
    assert rows["C4"][2] == ""                                    # not in SAGE
    for df in (win._display_df, win._omr_dl, win._octr_dl):
        assert {"Days (MR)", "Days (CTR)", "Local cost"} <= set(df.columns)


def test_the_total_box_uses_the_statistics_formula_and_updates_when_days_change():
    win = _window_with_costs(days={"C1": 3.0})
    win._local_mult_chk.setChecked(True)
    H.pump(5)
    names = _names(win._tab_matched)
    assert names[-2:] == ["Local cost total (MR)", "Local cost total (CTR)"] and "Local cost" not in names
    by = {c: (a, b) for c, a, b in zip(_col(win._tab_matched, "Stock Code"), _col(win._tab_matched, "Local cost total (MR)"),
                                       _col(win._tab_matched, "Local cost total (CTR)"))}
    assert by["C1"] == ("150.0", "180.0")                          # 3 days x 10 x qty 5 (MR) / qty 6 (CTR)
    assert by["C3"] == ("210.0", "210.0")                          # consumable: 30 x 7, no days
    row_c2 = _col(win._tab_matched, "Stock Code").index("C2")
    assert _col(win._tab_matched, "Local cost total (MR)")[row_c2] == ""                  # equipment without days
    assert "days are missing" in (_tip(win._tab_matched, row_c2, "Local cost total (MR)") or "")   # ... and the cell says why
    win._stats_panel.set_chosen_days("C2", 2.0)                    # the user types days in the Statistics tab
    H.pump(15)
    assert _col(win._tab_matched, "Local cost total (CTR)")[row_c2] == "360.0"            # 2 x 20 x qty 9
    win._local_mult_chk.setChecked(False)
    H.pump(5)
    assert _names(win._tab_matched)[-1] == "Local cost" and "Local cost total (MR)" not in _names(win._tab_matched)


def test_every_empty_days_cell_says_why_and_the_missing_ones_are_orange():
    import app as A
    df = pd.DataFrame({"Stock Code": ["C1", "C3", "C2", "X9"], "Qty (MR)": [1, 1, 1, 1], "Qty (CTR)": [None] * 4,
                       "MR Document": ["mr", "mr", "mr", "mr"], "MR Sheet": "S", "CTR Document": "", "CTR Sheet": ""})
    ctx = _ctx(COSTS, {("MR", "mr", "S", "C1"): (cs.EQUIPMENT, float("nan"), ""),          # equipment, no days: orange
                       ("MR", "mr", "S", "C3"): (cs.CONSUMABLE, float("nan"), ""),         # consumable: not billed per day
                       ("MR", "mr", "S", "C2"): ("", float("nan"), "unknown"),            # type not decided
                       ("MR", "mr", "S", "X9"): (cs.EQUIPMENT, 5.0, "")})                 # has days: nothing to explain
    _out, tips = A.MainWindow._with_local_costs(df, ctx, total=True)
    days_tip = {r: t for (r, c), t in tips.items() if c == "Days (MR)"}
    assert days_tip[0].startswith("Equipment needs Days")
    assert "consumable is not billed per day" in days_tip[1]
    assert "not decided" in days_tip[2]
    assert 3 not in days_tip
    ctr_tip = {r: t for (r, c), t in tips.items() if c == "Days (CTR)"}
    assert all("No CTR line" in t for t in ctr_tip.values()) and len(ctr_tip) == 4          # no CTR line on any of them


def test_the_missing_days_cells_are_painted_orange_in_the_tables_and_say_why_on_hover():
    win = _window_with_costs(days={"C1": 3.0})                  # C2 is equipment without days, C3 a consumable
    view = win._tab_matched
    names = _names(view)
    rows = _col(view, "Stock Code")
    r2, r3, r1 = rows.index("C2"), rows.index("C3"), rows.index("C1")
    m = view.model()

    def bg(row, col):
        c = m.index(row, names.index(col))
        b = m.data(c, Qt.ItemDataRole.BackgroundRole)
        return b.color().name() if b is not None else None
    assert bg(r2, "Days (MR)") == "#ffe0b2" and bg(r2, "Days (CTR)") == "#ffe0b2"           # missing days: orange
    assert bg(r1, "Days (MR)") != "#ffe0b2" and bg(r3, "Days (MR)") != "#ffe0b2"            # given / not applicable: not
    assert "needs Days" in _tip(view, r2, "Days (MR)")
    assert "consumable is not billed per day" in _tip(view, r3, "Days (MR)")
    win._stats_panel.set_chosen_days("C2", 2.0)                                               # once the days are typed ...
    H.pump(15)
    assert bg(_col(view, "Stock Code").index("C2"), "Days (MR)") != "#ffe0b2"                # ... the orange goes


def test_the_statistics_cost_view_shows_days_too_with_a_reason_for_the_empty_ones():
    import stats.stats_view as SV
    win = _window_with_costs()                                   # no days typed: the equipment lines wait
    p = win._stats_panel
    p.set_mode(SV.MODE_COST)
    H.pump(5)
    assert not p._tree.isColumnHidden(SV.C_DAYS)
    seen = {}
    tree = p._tree
    for i in range(tree.topLevelItemCount()):
        g = tree.topLevelItem(i)
        for j in range(g.childCount()):
            d = g.child(j)
            d.setExpanded(True)
            for n in range(d.childCount()):
                it = d.child(n)
                seen.setdefault(it.text(0), it)
    c1, c3 = seen["C1"], seen["C3"]
    assert c1.text(SV.C_DAYS) == "type days" and c1.background(SV.C_DAYS).color().name() == "#ffe0b2"
    assert "days" in c1.toolTip(SV.C_DAYS).lower()
    assert c3.text(SV.C_DAYS) == "—" and "not billed per day" in c3.toolTip(SV.C_DAYS)
    p.set_chosen_days("C1", 4.0)
    H.pump(15)
    assert float(next(it for it in _items_of(p) if it.text(0) == "C1" and it.text(SV.C_DAYS)[0].isdigit()).text(SV.C_TOTAL).replace(",", "")) > 0


def _items_of(panel):
    out = []
    tree = panel._tree
    for i in range(tree.topLevelItemCount()):
        g = tree.topLevelItem(i)
        for j in range(g.childCount()):
            d = g.child(j)
            d.setExpanded(True)
            for n in range(d.childCount()):
                out.append(d.child(n))
    return out


def test_the_box_only_shows_on_matched_and_the_two_only_tabs():
    win = _window_with_costs()
    for stack in (win._stack_matched, win._stack_only_mr, win._stack_only_ctr):
        win._tabs.setCurrentWidget(stack)
        H.pump(3)
        assert win._local_mult_chk.isVisibleTo(win), stack
    for tab in (win._stack_review, win._stack_error, win._stats_tab):
        win._tabs.setCurrentWidget(tab)
        H.pump(3)
        assert not win._local_mult_chk.isVisibleTo(win), tab


def test_sorting_and_hidden_columns_survive_the_refresh():
    win = _window_with_costs(days={"C1": 3.0})
    view = win._tab_matched
    view.sortByColumn(_names(view).index("Stock Code"), Qt.SortOrder.DescendingOrder)
    view.setColumnHidden(1, True)                          # hide Description (MR)
    win._local_mult_chk.setChecked(True)
    H.pump(5)
    codes = _col(view, "Stock Code")
    assert codes == sorted(codes, reverse=True)            # still sorted
    assert view.isColumnHidden(1)                          # still hidden


def test_the_excel_report_has_days_and_local_cost_at_the_end_of_the_three_sheets():
    import openpyxl
    from PySide6.QtWidgets import QFileDialog
    win = _window_with_costs(days={"C1": 3.0})
    out = H.SANDBOX / "local_cost_report.xlsx"
    real = QFileDialog.getSaveFileName
    QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(out), ""))
    try:
        win._download_report()
    finally:
        QFileDialog.getSaveFileName = real
    wb = openpyxl.load_workbook(out, data_only=True)
    for sheet in ("Matched", "Only in MR", "Only in CTR"):
        head = [c.value for c in wb[sheet][1]]
        if head and head[0] is not None:
            assert head[-3:] == ["Days (MR)", "Days (CTR)", "Local cost"], (sheet, head)
    matched = {r[0]: r[-1] for r in wb["Matched"].iter_rows(min_row=2, values_only=True)}
    assert matched["C1"] == 10.0 and matched["C3"] == 30.0 and matched["C4"] is None


# ── the Combined View table ──────────────────────────────────────────────────

COMBINED_COSTS = {"C1": 10.0, "C2": 20.0, "C3": 30.0}


def _combined_window(typed_days=None):
    """C1: MR 2 + CTR 3 (days 4); C2 (equipment): MR in two documents (5 and 1), CTR in two CTRs (4 x 7 days,
    5 x 182 days); C3 consumable: MR 7, CTR 7."""
    import app as A
    ST.save_types({"equipment": {"C1", "C2"}, "consumable": {"C3"},
                   "confirmed": {"C1": "Equipment", "C2": "Equipment", "C3": "Consumable"}})
    win = A.MainWindow()
    win.resize(1400, 800)
    win.show()
    H.pump(10)
    win._mr_df = pd.DataFrame({
        "Stock Code": ["C1", "C2", "C2", "C3"], "Qty": [2, 5, 1, 7], "Unit": "EA", "Description": "mr",
        "_SourceFile": ["mr1.xlsx", "mr1.xlsx", "mr2.xlsx", "mr1.xlsx"], "_SourceSheet": "S"})
    win._ctr_df = pd.DataFrame({
        "Stock Code": ["C1", "C2", "C2", "C3"], "Quantity": [3, 4, 5, 7], "Unit": "EA", "Description": "ctr",
        "Rate": "1", "Days": [4, 7, 182, None], "_SourceFile": ["a.xlsx", "a.xlsx", "b.xlsx", "a.xlsx"],
        "_SourceSheet": "P", "_Section": "Plant & Equipment"})
    win._run_compare()
    H.pump(30)
    assert H.wait(lambda: not win._stats_workers, 60)
    win._stats_panel.set_rows(cs.cost_rows(win._stats_items, COMBINED_COSTS, True))
    if typed_days:
        win._stats_panel.set_days(typed_days)
    win._local_ready = True
    win._apply_local_costs()
    win._combined_btn.setChecked(True)                   # Combined View
    H.pump(15)
    return win


def test_the_combined_table_ends_with_days_and_local_cost():
    win = _combined_window({"C2": 7.0})
    names = _names(win._tab_matched)
    assert names[-3:] == ["Days (MR)", "Days (CTR)", "Local cost"] and names.index("Flag") < names.index("Days (MR)")
    rows = dict(zip(_col(win._tab_matched, "Stock Code"), zip(_col(win._tab_matched, "Days (MR)"),
                                                              _col(win._tab_matched, "Days (CTR)"),
                                                              _col(win._tab_matched, "Local cost"))))
    assert rows["C1"] == ("4.0", "4.0", "10.0")                  # MR days taken from the matching CTR
    assert rows["C2"][1] == "7 / 182"                            # the two CTRs disagree: both shown
    assert rows["C3"] == ("", "", "30.0")                        # a consumable: no days


def test_the_combined_total_is_summed_line_by_line_because_the_lines_carry_different_days():
    win = _combined_window({"C2": 7.0})                          # the user picked 7 days for the MR lines of C2
    win._local_mult_chk.setChecked(True)
    H.pump(10)
    names = _names(win._tab_matched)
    assert names[-2:] == ["Local cost total (MR)", "Local cost total (CTR)"]
    tot = {c: (a, b) for c, a, b in zip(_col(win._tab_matched, "Stock Code"), _col(win._tab_matched, "Local cost total (MR)"),
                                        _col(win._tab_matched, "Local cost total (CTR)"))}
    assert tot["C1"] == ("80.0", "120.0")                        # 4 days x 10 x 2 (MR) / x 3 (CTR)
    assert tot["C2"] == ("840.0", "18760.0")                     # MR: 20 x (5+1) x 7; CTR: 20 x 4 x 7 + 20 x 5 x 182
    assert tot["C3"] == ("210.0", "210.0")                       # consumable: cost x quantity
    row = _col(win._tab_matched, "Stock Code").index("C2")
    assert "7 days" in _tip(win._tab_matched, row, "Days (CTR)") and "182 days" in _tip(win._tab_matched, row, "Days (CTR)")


def test_a_combined_total_without_days_is_blank_orange_and_says_why():
    win = _combined_window()                                     # nothing typed: the MR lines of C2 have no days
    win._local_mult_chk.setChecked(True)
    H.pump(10)
    view = win._tab_matched
    names = _names(view)
    row = _col(view, "Stock Code").index("C2")
    assert _col(view, "Local cost total (MR)")[row] == ""
    assert "days are missing" in _tip(view, row, "Local cost total (MR)")
    assert "needs Days" in _tip(view, row, "Days (MR)")
    m = view.model()
    b = m.data(m.index(row, names.index("Days (MR)")), Qt.ItemDataRole.BackgroundRole)
    assert b is not None and b.color().name() == "#ffe0b2"       # the missing Days cell is orange here too
    win._stats_panel.set_chosen_days("C2", 7.0)                  # typing them in the Statistics tab fills it in
    H.pump(15)
    assert _col(view, "Local cost total (MR)")[_col(view, "Stock Code").index("C2")] == "840.0"


def test_switching_combined_view_off_and_on_keeps_the_columns():
    win = _combined_window({"C2": 7.0})
    win._combined_btn.setChecked(False)
    H.pump(10)
    assert _names(win._tab_matched)[-3:] == ["Days (MR)", "Days (CTR)", "Local cost"] and "Flag" not in _names(win._tab_matched)
    win._combined_btn.setChecked(True)
    H.pump(10)
    assert _names(win._tab_matched)[-3:] == ["Days (MR)", "Days (CTR)", "Local cost"] and "Flag" in _names(win._tab_matched)


def test_the_excel_combined_sheet_has_the_columns_too():
    import openpyxl
    from PySide6.QtWidgets import QFileDialog
    win = _combined_window({"C2": 7.0})
    out = H.SANDBOX / "combined_report.xlsx"
    real = QFileDialog.getSaveFileName
    QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(out), ""))
    try:
        win._download_report()
    finally:
        QFileDialog.getSaveFileName = real
    ws = openpyxl.load_workbook(out, data_only=True)["Combined"]
    head = [c.value for c in ws[1]]
    assert head[-3:] == ["Days (MR)", "Days (CTR)", "Local cost"], head
    got = {r[0]: r[-1] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert got["C2"] == 20.0 and got["C3"] == 30.0


def test_a_new_compare_does_not_show_the_previous_compares_costs():
    win = _combined_window({"C2": 7.0})
    win._mr_df = win._mr_df.copy()
    win._run_compare()                                           # compare again: the statistics start from scratch
    H.pump(10)
    assert set(_col(win._tab_matched, "Local cost")) <= {"", "⏳ waiting"}   # no figure until THIS compare's statistics are ready
    assert win._combined_df["Local cost"].isna().all()
    assert H.wait(lambda: not win._stats_workers, 60)

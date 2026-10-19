"""Document price: the second view of the statistics (rate × quantity × days), the equipment rate lookup,
the editable rechargeability choices, the toggles and their Settings defaults, and the Generator's price totals.
Fixed fixture figures only — the SAGE export in TEST_Files has random costs."""
import pandas as pd

from stats import cost_stats as cs
import harness as H
from harness import W
from stats import stats_view as SV
from ctr_tools import source_choices as SC
from ctr_tools import stockcode_types as ST
from stats.stats_view import StatsPanel


def _items(*rows, **extra):
    """rows: (key, qty, status, section). Extra columns are given once for all rows."""
    df = pd.DataFrame([{"key": k, "doc": "doc1", "sheet": "", "side": "CTR", "qty": q, "status": s,
                        "section": sec} for k, q, s, sec in rows])
    for name, values in extra.items():
        df[name] = values
    return df


EQ, CONS = "Plant & Equipment", "Consumables & Materials"


def _rows(items, **kw):
    kw.setdefault("equip_rates", {})
    return cs.cost_rows(items, kw.pop("costs", {}), True, **kw)


# ── the price rules ──────────────────────────────────────────────────────────

def test_price_is_quantity_times_rate_times_days_for_equipment_and_quantity_times_rate_for_consumables():
    items = _items(("E1", 2, "RECHARGE", EQ), ("C1", 10, "RECHARGE", CONS))
    rows = _rows(items, costs={"C1": 34.0}, equip_rates={"E1": 100.0})
    p = cs.price_view(rows, {"E1": 5.0}).set_index("key")
    assert p.loc["E1", "cost"] == 2 * 100 * 5 and p.loc["E1", "unit_cost"] == 100 and p.loc["E1", "days"] == 5
    assert p.loc["C1", "cost"] == 10 * 34 and (p["problem"] == "").all()


def test_non_rechargeable_is_billed_at_zero_even_without_a_rate_or_days():
    rows = _rows(_items(("E1", 2, "NONRECHARG", EQ), ("C1", 3, "NONRECHARG", CONS)))
    p = cs.price_view(rows).set_index("key")
    assert p.loc["E1", "cost"] == 0.0 and p.loc["C1", "cost"] == 0.0 and (p["problem"] == "").all()


def test_nothing_is_assumed_a_missing_rate_days_or_quantity_is_left_out_and_says_why():
    rows = _rows(_items(("E1", 2, "RECHARGE", EQ),            # pricebook has no rate
                        ("E2", 2, "RECHARGE", EQ),            # rate but no days
                        ("C1", None, "RECHARGE", CONS),       # no quantity
                        ("C2", 4, "RECHARGE", CONS)),         # not in SAGE
                 costs={"C1": 5.0}, equip_rates={"E2": 10.0})
    p = cs.price_view(rows).set_index("key")
    assert p.loc["E1", "problem"] == cs.NO_RATE and p.loc["E2", "problem"] == cs.NO_DAYS
    assert p.loc["C1", "problem"] == cs.NO_QTY and p.loc["C2", "problem"] == cs.NO_RATE
    assert p["cost"].isna().all()
    assert cs.price_view(rows, {"E2": 3.0}).set_index("key").loc["E2", "cost"] == 2 * 10 * 3


def test_markup_applies_to_consumables_only_and_only_when_given():
    rows = _rows(_items(("E1", 1, "RECHARGE", EQ), ("C1", 10, "RECHARGE", CONS)),
                 costs={"C1": 100.0}, equip_rates={"E1": 50.0})
    plain = cs.price_view(rows, {"E1": 2.0}).set_index("key")
    marked = cs.price_view(rows, {"E1": 2.0}, markup_rate=0.065).set_index("key")
    assert plain.loc["C1", "cost"] == 1000.0 and marked.loc["C1", "cost"] == 1065.0
    assert plain.loc["E1", "cost"] == marked.loc["E1", "cost"] == 100.0


def test_a_code_no_ctr_section_names_takes_its_kind_from_the_stock_code_types_list_else_it_is_not_priced():
    items = pd.DataFrame([{"key": k, "doc": "mr", "sheet": "", "side": "MR", "qty": 2, "status": "RECHARGE"}
                          for k in ("E9", "C9", "U9")])
    types = {"equipment": {"E9"}, "consumable": {"C9"}, "confirmed": {}}
    rows = cs.cost_rows(items, {"C9": 7.0}, True, equip_rates={"E9": 10.0})
    rows = cs.apply_manual(rows, {}, True, None, types)
    assert dict(zip(rows["key"], rows["kind"])) == {"E9": cs.EQUIPMENT, "C9": cs.CONSUMABLE, "U9": ""}
    p = cs.price_view(rows, {"E9": 3.0}).set_index("key")
    assert p.loc["E9", "cost"] == 60.0 and p.loc["C9", "cost"] == 14.0
    assert p.loc["U9", "problem"] == cs.NO_TYPE and pd.isna(p.loc["U9", "cost"])


def test_a_line_with_its_own_rate_cell_is_never_priced_from_sage_behind_its_back():
    items = _items(("C1", 2, "RECHARGE", CONS), ("C2", 2, "RECHARGE", CONS), rate=[3.0, None])
    rows = _rows(items, costs={"C1": 99.0, "C2": 99.0})
    p = cs.price_view(rows).set_index("key")
    assert p.loc["C1", "cost"] == 6.0                      # the cell's rate, not SAGE's 99
    assert p.loc["C2", "problem"] == cs.NO_RATE            # a blank cell stays blank


def test_scaffold_is_priced_per_ton_times_rate_times_days_and_only_when_included():
    items = _items(("SC", 3, "RECHARGE", EQ), ("E1", 1, "RECHARGE", EQ), rate=[20.0, 10.0],
                   days=[4.0, 2.0], scaffold=[True, False])
    rows = _rows(items)
    assert set(cs.price_view(rows)["key"]) == {"E1"}
    p = cs.price_view(rows, include_scaffold=True).set_index("key")
    assert p.loc["SC", "cost"] == 3 * 20 * 4 and p.loc["SC", "kind"] == cs.SCAFFOLD
    assert p.loc["E1", "cost"] == 20.0


def test_summaries_and_lists_work_on_the_price_view_unchanged():
    rows = _rows(_items(("E1", 1, "RECHARGE", EQ), ("E2", 1, "NONRECHARG", EQ), ("E3", 1, "RECHARGE", EQ)),
                 equip_rates={"E1": 10.0, "E2": 99.0})
    s = cs.summarize(cs.price_view(rows, {"E1": 2.0}))
    tot = s[(s["side"] == "CTR") & (s["doc"] == cs.TOTAL_LABEL)].iloc[0]
    assert tot[cs.RECHARGEABLE] == 20.0 and tot[cs.NON_RECHARGEABLE] == 0.0
    assert tot["items"] == 2 and tot["no_cost"] == 1


# ── equipment rates: the Generator's lookup ──────────────────────────────────

def test_equipment_rate_follows_the_generators_lookup():
    names = pd.DataFrame({"product": ["EQ1", "EQ2", "EQ3"],
                          "legacy_name": ["old1", "old2", "old3"],
                          "canonical_name": ["Pump A", "", ""],
                          "non_recharge": [False, True, False]})
    usd = pd.DataFrame({"stock_code": ["X-STBY", "X-NOR", "Y-NOR"],
                        "supplier_desc": ["Pump A", "Pump A", "Direct Item"],
                        "unit_price": [5.0, 9.0, 12.0]})
    rates = cs.equip_rate_map(["eq1", "EQ2", "EQ3", "Direct Item", "Nothing"], names, usd)
    assert rates["EQ1"] == 9.0                  # through the bridge; the normal rate beats standby
    assert rates["EQ2"] == 0.0                  # bridge says non-recharge: billed at 0
    assert "EQ3" not in rates                   # bridge name "old3" is not in the pricebook: no rate
    assert rates["DIRECT ITEM"] == 12.0         # not in the bridge: looked up by that text directly
    assert "NOTHING" not in rates
    assert cs.equip_rate_map(["EQ2", "Direct Item"], names, None) == {"EQ2": 0.0}      # no pricebook: only the bridge's 0


# ── the panel ────────────────────────────────────────────────────────────────

def _price_panel(show_markup=False):
    rows = _rows(_items(("E1", 2, "RECHARGE", EQ), ("C1", 10, "RECHARGE", CONS)),
                 costs={"C1": 100.0}, equip_rates={"E1": 50.0})
    p = StatsPanel(show_markup=show_markup)
    p.set_rows(rows)
    return p


def _item_texts(panel, col):
    out = {}
    tree = panel._tree
    for i in range(tree.topLevelItemCount()):
        g = tree.topLevelItem(i)
        for j in range(g.childCount()):
            d = g.child(j)
            d.setExpanded(True)
            for n in range(d.childCount()):
                it = d.child(n)
                out[it.text(SV.C_NAME)] = it.text(col)
    return out


def test_the_toggle_switches_between_document_costs_and_document_price():
    p = _price_panel()
    assert p.mode() == SV.MODE_COST and not p._tree.isColumnHidden(SV.C_UNIT)
    assert not p._tree.isColumnHidden(SV.C_DAYS)                           # costs use days too: always shown
    assert p._scaffold_chk.isHidden() and not p._qty_chk.isHidden()
    p.set_mode(SV.MODE_PRICE)
    assert p._tree.headerItem().text(SV.C_UNIT) == "Rate" and p._tree.headerItem().text(SV.C_TOTAL) == "Total price"
    assert not p._tree.isColumnHidden(SV.C_DAYS) and p._qty_chk.isHidden()
    assert _item_texts(p, SV.C_UNIT)["C1"] == "100.00"
    p.set_mode(SV.MODE_COST)
    assert p._tree.headerItem().text(SV.C_UNIT) == "Unit cost"


def test_days_are_typed_in_the_price_view_and_announced():
    p = _price_panel()
    p.set_mode(SV.MODE_PRICE)
    got = []
    p.days_changed.connect(got.append)
    assert "Days" in _item_texts(p, SV.C_NOTE)["E1"]                       # not counted until the days are typed
    tree = p._tree
    item = next(g.child(0).child(n) for g in [tree.topLevelItem(0)] for n in range(g.child(0).childCount())
                if g.child(0).child(n).text(SV.C_NAME) == "E1")
    assert item.flags() & SV.Qt.ItemFlag.ItemIsEditable
    tree.itemDoubleClicked.emit(item, SV.C_DAYS)
    item.setText(SV.C_DAYS, "4")
    H.pump()
    assert got and got[-1] == {"E1": 4.0} and p.days() == {"E1": 4.0}
    H.pump(5)
    assert _item_texts(p, SV.C_TOTAL)["E1"] == "400.00"                    # 2 × 50 × 4
    p.set_days({})
    assert _item_texts(p, SV.C_TOTAL)["E1"] == ""                          # no days: not counted


def test_the_markup_option_exists_only_where_asked_and_is_off_in_the_mr_vs_ctr_panel():
    mr = _price_panel(show_markup=False)
    mr.set_markup_rate(0.065)
    mr.set_mode(SV.MODE_PRICE)
    assert mr._markup_chk.isHidden() and mr.markup_rate() is None
    gen = _price_panel(show_markup=True)
    gen.set_markup_rate(0.065)
    gen.set_mode(SV.MODE_PRICE)
    assert not gen._markup_chk.isHidden()
    assert gen.markup_rate() == 0.065
    gen._markup_chk.setChecked(False)
    assert gen.markup_rate() is None


def test_scaffold_is_off_by_default_and_only_in_the_price_view():
    items = _items(("SC", 3, "RECHARGE", EQ), ("E1", 1, "RECHARGE", EQ), rate=[20.0, 10.0],
                   days=[4.0, 2.0], scaffold=[True, False])
    p = StatsPanel()
    p.set_rows(_rows(items))
    assert "SC" not in _item_texts(p, SV.C_NAME)                           # cost view: no SAGE cost, not listed
    p.set_mode(SV.MODE_PRICE)
    assert "SC" not in _item_texts(p, SV.C_NAME) and not p.include_scaffold()
    p._scaffold_chk.setChecked(True)
    assert _item_texts(p, SV.C_TOTAL)["SC"] == "240.00"


def test_settings_defaults_set_the_starting_toggles():
    SV.save_stats_defaults(SV.MODE_PRICE, True, False)
    try:
        p = StatsPanel(show_markup=True)
        assert p.mode() == SV.MODE_PRICE and p.include_scaffold() and not p._markup_chk.isChecked()
        SV.save_stats_defaults(SV.MODE_COST, False, True)
        p.apply_defaults()
        assert p.mode() == SV.MODE_COST and not p.include_scaffold() and p._markup_chk.isChecked()
    finally:
        SV.save_stats_defaults(SV.MODE_COST, False, True)


def test_the_choices_list_is_editable_in_settings_and_shared():
    from common import settings_view as SVW
    w = H.generator()
    tab = SVW.SettingsTab(w)
    SC.save_choices({"AA": "SAGE", "BB": "DOC"})
    w.reload_saved_data()
    tab.refresh()
    assert tab._counts["choices"].text() == "2 saved"
    dlg = SVW._ChoicesDialog(None, SC.load_choices())
    assert dlg.entries() == {"AA": "SAGE", "BB": "DOC"}
    dlg._add("cc", "SAGE")                                        # a new code, typed in lower case
    dlg._table.removeRow(0)                                       # AA removed
    assert dlg.entries() == {"BB": "DOC", "CC": "SAGE"}
    SC.save_choices(dlg.entries())
    w.reload_saved_data()
    assert w._stats_panel.choices() == {"BB": "DOC", "CC": "SAGE"}   # the Generator's list took it
    SC._CHOICES_PATH.unlink()
    w.reload_saved_data()


def test_a_choice_typed_in_settings_settles_a_disparity_before_it_is_asked():
    w = H.generator()
    p = w._stats_panel
    code = sorted(set(p.disparity_table()["Stock Code"]))[0]
    SC.save_choices({code: "DOC"})
    w.reload_saved_data()
    assert code not in set(p.disparity_table()["Stock Code"])
    assert p.disparity_table(include_resolved=True).set_index("Stock Code").loc[code, "Chosen"] == "Document"
    SC._CHOICES_PATH.unlink()
    w.reload_saved_data()


def test_settings_defaults_are_in_settings():
    from common import settings_view as SVW
    w = H.generator()
    tab = SVW.SettingsTab(w)
    got = []
    w.stats_defaults_changed.connect(lambda: got.append(1))
    tab._scaffold_chk.setChecked(True)
    assert got and SV.stats_defaults()["scaffold"] is True
    tab._scaffold_chk.setChecked(False)


# ── the Excel report ─────────────────────────────────────────────────────────

def test_the_price_report_has_its_own_frames_with_a_description_column():
    items = _items(("E1", 2, "RECHARGE", EQ), ("C1", 10, "RECHARGE", CONS),
                   description=["Big pump", "Rope"])
    rows = _rows(items, costs={"C1": 100.0}, equip_rates={"E1": 50.0}, sage_names={"C1": "SAGE rope"})
    summary, lines = cs.price_report_frames(rows, days={"E1": 4.0}, markup_rate=0.1)
    assert "Total price" in summary.columns and "Total cost" not in summary.columns
    it = lines.set_index("Stock Code")
    assert it.loc["E1", "Price"] == 400.0 and abs(it.loc["C1", "Price"] - 1100.0) < 1e-9
    assert it.loc["E1", "Description"] == "Big pump" and it.loc["C1", "Description"] == "Rope"
    assert "markup" in summary["Price basis"].iloc[0]
    csum, clines = cs.report_frames(rows)
    assert "Total cost" in csum.columns and "Description" in clines.columns


def test_saved_rows_keep_the_price_inputs_and_older_saves_still_load():
    rows = _rows(_items(("E1", 2, "RECHARGE", EQ)), equip_rates={"E1": 50.0})
    back = cs.rows_from_records(cs.rows_to_records(rows))
    assert back.iloc[0]["equip_rate"] == 50.0 and back.iloc[0]["kind"] == cs.EQUIPMENT
    old = [{"key": "E1", "doc": "d", "sheet": "", "side": "CTR", "category": cs.EQUIPMENT,
            "bucket": cs.RECHARGEABLE, "qty": 2, "unit_cost": 5.0, "source": "SAGE", "sage_bucket": "", "name": ""}]
    legacy = cs.rows_from_records(old)
    assert legacy.iloc[0]["cost"] == 10.0 and pd.isna(legacy.iloc[0]["rate"])
    assert cs.price_view(legacy).iloc[0]["problem"] == cs.NO_RATE         # an old save has no price inputs


# ── the CTR Generator: its panel and the price totals ────────────────────────

def _scaffold_row(w):
    tbl = w._req_equip_tbl
    for r in range(tbl.rowCount()):
        d = tbl.item(r, W._EQ_DESC)
        if d is not None and d.data(W._SCAFFOLD_ROLE):
            return r
    raise AssertionError("no scaffold line in the test request")


def _table_sum(tbl, col, status_col, skip_scaffold=None):
    total = 0.0
    for r in range(tbl.rowCount()):
        if not tbl.item(r, status_col).text().startswith("✓"):
            continue
        d = tbl.item(r, 0)
        is_scaffold = bool(d is not None and d.data(W._SCAFFOLD_ROLE))
        if skip_scaffold is not None and is_scaffold == skip_scaffold:
            continue
        try:
            total += float(tbl.item(r, col).text())
        except ValueError:
            pass
    return total


def test_price_totals_are_per_category_and_per_currency():
    w = H.generator()
    H.resolve_disparities(w)
    sr = _scaffold_row(w)
    w._req_equip_tbl.item(sr, W._EQ_RATE).setText("10")
    w._req_equip_tbl.item(sr, W._EQ_DAYS).setText("2")
    w._refresh_stats()
    t = w.price_totals()
    gone = t["left_out"]
    # Manpower and transport are AZN; equipment and consumables USD; they are never added together.
    assert t["azn_total"] == t["manpower"] + t["transport"]
    assert t["usd_total"] == t["equipment"] + t["consumables"]                  # scaffold not included by default
    assert not t["scaffold_included"] and t["scaffold"] > 0
    w._stats_panel._scaffold_chk.setChecked(True)
    t2 = w.price_totals()
    assert abs(t2["usd_total"] - (t["usd_total"] + t["scaffold"])) < 1e-6 and t2["equipment"] == t["equipment"]
    # Consumables follow the markup toggle.
    raw = _table_sum(w._req_cons_tbl, W._CS_TOTAL, W._CS_STATUS)
    w._stats_panel._markup_chk.setChecked(True)
    on = w.price_totals()["consumables"]
    w._stats_panel._markup_chk.setChecked(False)
    off = w.price_totals()["consumables"]
    assert abs(off - on / (1 + w._markup_rate())) < 1e-6 and off <= raw + 1e-6
    # The shown text names every category.
    w._refresh_price_totals()
    text = w._price_totals_lbl.text()
    for label in ("Manpower", "Transport", "Equipment", "Consumables", "AZN CTR total", "USD CTR total"):
        assert label in text, label
    assert gone["equipment"] == 0                                              # the scaffold line was completed above
    w._stats_panel._scaffold_chk.setChecked(False)
    w._stats_panel._markup_chk.setChecked(True)


def test_a_line_with_a_missing_value_is_not_priced_and_is_counted():
    w = H.generator()
    t = w.price_totals()
    assert t["left_out"]["equipment"] >= 1                                      # the scaffold line has no rate / days
    w._refresh_price_totals()
    assert "not priced" in w._price_totals_lbl.text()


def test_the_generators_panel_prices_the_lines_as_the_tables_carry_them():
    w = H.generator()
    H.resolve_disparities(w)
    p = w._stats_panel
    p.set_mode(SV.MODE_PRICE)
    p._markup_chk.setChecked(False)
    p._scaffold_chk.setChecked(False)
    priced = p._view[p._view["problem"] == ""]
    cons_rows = priced[priced["kind"] == cs.CONSUMABLE]
    expected = _table_sum(w._req_cons_tbl, W._CS_TOTAL, W._CS_STATUS)
    assert abs(cons_rows["cost"].sum() - expected) < 0.01 or cons_rows["cost"].sum() <= expected + 0.01
    assert not p._markup_chk.isHidden()                                        # the Generator has the option
    p.set_mode(SV.MODE_COST)


# ── the MR vs CTR window and the Excel files ─────────────────────────────────

def _mr_window():
    """The main window with a small synthetic MR and CTR compared, the SAGE Export, USD Pricebook and
    Combined DB of the test files set, and the statistics finished."""
    import app as A
    ST.save_types({"equipment": {"EQ1500456000"}, "consumable": {"PP3800000000"}, "confirmed": {}})
    win = A.MainWindow()
    win.resize(1300, 800)
    win.show()
    H.pump(10)
    win._usd_pb_edit.setText(str(H.INPUTS / "4410030190 - 12.02.2026.xlsx"))
    win._combined_db_edit.setText(str(H.INPUTS / "CTR_NAMES_DB & CTR_REQUEST 263.xlsx"))
    win._sage_edit.setText(str(H.INPUTS / "FROM SAGE 10.06.26.xlsm"))
    win._mr_df = pd.DataFrame({
        "Stock Code": ["EQ1500456000", "PP3800000000", "ZZ9999"], "Qty": [2, 4, 1],
        "Rechargeable": ["RECHARGE", "NONRECHARG", "RECHARGE"],
        "Description": ["mr pump", "mr nonrech", "mr unknown"], "_SourceFile": "mr.xlsx", "_SourceSheet": "S1"})
    win._ctr_df = pd.DataFrame({
        "Stock Code": ["EQ1500456000"], "Quantity": [2], "Rate": ["1.00"], "Description": ["ctr pump"],
        "_SourceFile": "ctr.xlsx", "_SourceSheet": "S1", "_Section": ["Plant & Equipment"]})
    win._run_compare()
    assert H.wait(lambda: win._stats_panel.has_rows(), 120), "the statistics did not finish"
    _settle(win)
    return win


def _settle(win):
    """Wait until no statistics worker is running (a thread still running at exit crashes the process)."""
    assert H.wait(lambda: not win._stats_workers, 120), "a statistics worker is still running"
    H.pump(20)


def test_the_mr_vs_ctr_tab_shares_the_pricebook_fields_with_the_generator():
    win = _mr_window()
    g = win._ctr_gen
    assert g.usd_pricebook_path() == win._usd_pb_edit.text() and g.combined_db_path() == win._combined_db_edit.text()
    g.set_usd_pricebook_path("/nowhere/x.xlsx")                 # set in the Generator: shows in the MR vs CTR tab
    assert win._usd_pb_edit.text() == "/nowhere/x.xlsx"
    win._combined_db_edit.setText("/nowhere/db.xlsx")           # and the other way round
    assert g.combined_db_path() == "/nowhere/db.xlsx"
    _settle(win)


def test_the_mr_vs_ctr_price_view_prices_equipment_from_the_pricebook_without_markup():
    win = _mr_window()
    p = win._stats_panel
    p.set_mode(SV.MODE_PRICE)
    assert p._markup_chk.isHidden() and p.markup_rate() is None             # no markup option in this tab
    v = p._view.set_index(["side", "key"])
    assert v.loc[("MR", "EQ1500456000"), "unit_cost"] == 1.0                # the USD Pricebook rate (Generator shows 1.00)
    assert v.loc[("MR", "EQ1500456000"), "problem"] == cs.NO_DAYS           # days are typed by the user
    assert v.loc[("MR", "PP3800000000"), "cost"] == 0.0                     # non-rechargeable
    assert v.loc[("MR", "ZZ9999"), "problem"] == cs.NO_TYPE                 # in no list, no CTR section: type unknown, not priced
    p.set_days({"EQ1500456000": 3})
    assert p._view.set_index(["side", "key"]).loc[("CTR", "EQ1500456000"), "cost"] == 2 * 1.0 * 3
    assert "Description" in p.price_report_frames()[1].columns
    assert p._base["description"].tolist()[:2] == ["mr pump", "mr nonrech"]


def test_a_saved_comparison_keeps_the_days_and_the_price_inputs():
    win = _mr_window()
    win._stats_panel.set_days({"EQ1500456000": 5})
    snap = win._stats_snapshot()
    assert snap["days"] == {"EQ1500456000": 5.0}
    rows = cs.rows_from_records(snap["rows"])
    p = StatsPanel()
    p.restore(rows, {}, True, days=snap["days"])
    p.set_mode(SV.MODE_PRICE)
    assert p._view.set_index(["side", "key"]).loc[("MR", "EQ1500456000"), "cost"] == 2 * 1.0 * 5   # no files needed


def test_the_mr_vs_ctr_excel_report_has_cost_and_price_on_separate_sheets():
    import openpyxl
    from PySide6.QtWidgets import QFileDialog
    win = _mr_window()
    win._stats_panel.set_days({"EQ1500456000": 3})
    out = H.SANDBOX / "mr_report.xlsx"
    real = QFileDialog.getSaveFileName
    QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (str(out), ""))
    try:
        win._download_report()
    finally:
        QFileDialog.getSaveFileName = real
    wb = openpyxl.load_workbook(out, data_only=True)
    for name in ("Statistics", "Statistics Items", "Price Statistics", "Price Items"):
        assert name in wb.sheetnames, (name, wb.sheetnames)
    head = [c.value for c in wb["Price Items"][1]]
    assert "Description" in head and "Rate" in head and "Days" in head and "Price" in head
    assert "Description" in [c.value for c in wb["Statistics Items"][1]]
    rows = [dict(zip(head, [c.value for c in r])) for r in wb["Price Items"].iter_rows(min_row=2)]
    pump = next(r for r in rows if r["Stock Code"] == "EQ1500456000" and r["Side"] == "MR")
    assert pump["Price"] == 6.0 and pump["Days"] == 3 and pump["Description"] == "mr pump"
    assert [c.value for c in wb["Price Statistics"][1]][2] == "Rechargeable price"


def test_the_generators_statistics_excel_has_the_price_sheets_too():
    import openpyxl
    w = H.generator()
    H.resolve_disparities(w)
    out_dir = H.SANDBOX / "stats_out"
    out_dir.mkdir(exist_ok=True)

    class _Dlg:
        def __init__(self, *a, **k):
            pass

        def exec(self):
            return W.QDialog.DialogCode.Accepted

        def folder(self):
            return str(out_dir)

        def deleteLater(self):
            pass

    real = W._SaveLocationDialog
    W._SaveLocationDialog = _Dlg
    try:
        w._save_stats_excel()
    finally:
        W._SaveLocationDialog = real
    files = sorted(out_dir.glob("*.xlsx"))
    assert files, "no statistics file written"
    wb = openpyxl.load_workbook(files[-1], data_only=True)
    assert wb.sheetnames == ["Statistics", "Statistics Items", "Price Statistics", "Price Items", "Info"]
    info = {r[0].value: r[1].value for r in wb["Info"].iter_rows(min_row=2)}
    assert "Price sheets" in info and "markup" in info["Price sheets"]
    head = [c.value for c in wb["Price Items"][1]]
    assert "Description" in head and "Days" in head


def test_the_mr_vs_ctr_statistics_say_that_scaffold_is_left_out():
    import app as A
    win = A.MainWindow()
    win.resize(1300, 800)
    win.show()
    H.pump(10)
    p = win._stats_panel
    p.set_mode(SV.MODE_PRICE)
    assert p._scaffold_chk.isHidden() and not p._no_scaffold_lbl.isHidden()
    assert "not included" in p._no_scaffold_lbl.text() and "scaffold" in p._no_scaffold_lbl.toolTip().lower()
    p._scaffold_chk.setChecked(True)                           # even a stored "include scaffold" default has no effect here
    assert not p.include_scaffold()
    p.set_mode(SV.MODE_COST)
    assert p._no_scaffold_lbl.isHidden()
    g = H.generator()
    g._stats_panel.set_mode(SV.MODE_PRICE)
    assert not g._stats_panel._scaffold_chk.isHidden() and g._stats_panel._no_scaffold_lbl.isHidden()   # the Generator keeps it

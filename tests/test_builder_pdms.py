"""PDMS proposal (builder_pdms): tags filled, tables built, optional blocks, totals, nothing left over.

Fixture rows only; the template under TEST_Files/CTR creator is read, the output goes to the sandbox.
"""
import docx

import harness as H
from ctr_tools import builder_pdms as B

TEMPLATE = H.INPUTS / "PDMS_TEMPLATE_short.docx"

HEADER = {
    "prp_number": "PRP-26-099", "date": "2026-09-09", "attention_name": "Jane Doe",
    "scope": "P000001 - Test job - REQ0000001", "project_type": "Offshore",
}
# As the Generator holds them: quantity in HOURS (60 days x 12 h) at an HOURLY rate (day rate / 12)
MAN = [
    {"comment": "Day Shift", "num_employees": 1, "description": "PA Level Foreman-National",
     "quantity": 720, "uom": "Hours", "rate_azn": 37.14, "nationality": "National", "hours_per_shift": 12},
    {"comment": "Day Shift", "num_employees": 2, "description": "Scaffolder-National",
     "quantity": 720, "uom": "Hours", "rate_azn": 22.53, "nationality": "National", "hours_per_shift": 12},
]
EQUIP = [
    {"description": "Conventional Scaffold", "quantity": 11.42, "unit": "TON", "rate_per_day": 2.5,
     "days": 62, "stock_code": "X", "sage_cost": 0.0},
    {"description": "Not billed", "quantity": 1, "unit": "EA", "rate_per_day": "NONRECHARG",
     "days": 5, "stock_code": "Y", "sage_cost": 0.0},
]
CONS = [
    {"long_description": "GLOVES, SIZE 9", "local_expect_cost": 2.0, "unit_code": "EA",
     "product": "P1", "quantity": 15, "sage_cost": 0.0},
]


def _build(tmp, equip=EQUIP, cons=CONS, header=HEADER, hotbed=None):
    return B.build_pdms(MAN, [], equip, cons, TEMPLATE, tmp, header, rates={"USD": 1.7}, markup_rate=0.1,
                        hotbed=hotbed)


def _text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def test_helpers():
    assert B.attention_from_requester("Jane Doe <a@x.com> ") == "Jane Doe"
    assert B.attention_from_requester("Doe, Jane (j@x.com) Maintenance Specialist") == "Jane Doe"
    assert B.salutation_from_attention("Jane Doe") == "Jane"
    assert B.letter_date("2026-09-09") == "9 September 2026"
    assert B.letter_date("soon") == "soon"


def test_full_proposal(tmp_path=None):
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")
    out = _build(H.SANDBOX / "pdms_full", hotbed={"rate": 59.0, "days": 10})
    assert out.name == "PRP-26-099 P000001 - Test job - REQ0000001.docx"
    d = docx.Document(str(out))
    text = _text(d)
    assert "{{" not in text and "}}" not in text
    assert "9 September 2026" in text and "Dear Jane," in text and "Attention: Jane Doe" in text
    cfg = B.pdms_defaults()                                  # the customer and signer come from the configuration
    assert cfg["customer_name"] and cfg["customer_name"] in text and cfg["signer_name"] in text
    assert "Offshore Work Schedule:" in text and "in AZN" in text
    assert "Rates are listed in Table № 1, Table № 2 and Table № 3." in text
    assert len(d.tables) == 4
    labour, equip, cons, summary = d.tables
    # the sample's own table: its two header rows, merged "LABOUR TOTAL:", fills and fonts
    assert [c.text.strip() for c in labour.rows[0].cells][4] == "No of"
    assert [c.text.strip() for c in labour.rows[1].cells][4] == "men"
    assert [c.text for c in labour.rows[2].cells][:2] == ["PA Level Foreman-Day shift", "National"]
    # hours become days (720 h / 12 = 60) and the hourly rate a day rate (37.14 x 12 = 445.68)
    assert [c.text for c in labour.rows[2].cells][3:7] == ["60", "1", "445.68 ₼", "26,740.80 ₼"]
    assert [c.text for c in labour.rows[3].cells][0] == "Scaffolder-Day shift"
    assert [c.text for c in labour.rows[3].cells][3:7] == ["60", "2", "270.36 ₼", "32,443.20 ₼"]
    # Hotbed: rate x days
    assert [c.text for c in labour.rows[4].cells][0] == "Hotbed – if required"
    assert [c.text for c in labour.rows[4].cells][3] == "10" and labour.rows[4].cells[5].text == "59.00 ₼"
    assert labour.rows[4].cells[6].text == "590.00 ₼"
    fills = lambda row: [c._tc.tcPr.find(docx.oxml.ns.qn("w:shd")).get(docx.oxml.ns.qn("w:fill")) for c in row.cells]
    assert fills(labour.rows[0])[0] == "D9D9D9" and fills(labour.rows[2])[0] == "DCE6F1" and fills(labour.rows[-1])[0] == "8EAADB"
    # 445.68*60*1 + 270.36*60*2 = 59 184.00, + hotbed 59 x 10
    assert labour.rows[-1].cells[-1].text == "59,774.00 ₼"
    # 11.42 * 62 * (2.5 * 1.7) = 3 009.17 ; the non-rechargeable line is not billed
    assert len(equip.rows) == 3 and equip.rows[0].cells[4].text == "Unit Price" and equip.rows[-1].cells[-1].text == "3,009.17 ₼"
    # 15 * (2.0 * 1.1 * 1.7) = 56.10
    assert cons.rows[-1].cells[-1].text == "56.10 ₼"
    grand = 59774 + 11.42 * 62 * 2.5 * 1.7 + 15 * 2.0 * 1.1 * 1.7
    assert summary.rows[-1].cells[-1].text == f"{grand:,.2f} ₼"
    assert f"{grand:,.2f}" in text
    for part in (d.sections[0].header, d.sections[0].first_page_header):
        assert [p.text for p in part.paragraphs if p.text.strip()][0].startswith("PRP-26-099")


def test_empty_sections_are_removed_and_tables_renumbered():
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")
    out = _build(H.SANDBOX / "pdms_labour", equip=[], cons=[])
    d = docx.Document(str(out))
    text = _text(d)
    assert "{{" not in text
    assert "Equipment rates" not in text and "Consumables costs" not in text
    assert "Rates are listed in Table № 1." in text
    assert len(d.tables) == 2                       # labour + summary
    assert [r.cells[0].text.strip() for r in d.tables[1].rows] == ["LABOUR:", "GRAND TOTAL:"]


def test_refuses_to_guess():
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")
    try:
        B.build_pdms(MAN, [], EQUIP, [], TEMPLATE, H.SANDBOX / "pdms_x", HEADER, rates={}, markup_rate=0)
    except ValueError as e:
        assert "No CBAR rate for USD" in str(e)
    else:
        raise AssertionError("a priced USD line without a USD rate must be refused")
    try:
        B.build_pdms(MAN, [], [], [], H.SANDBOX / "missing.docx", H.SANDBOX, HEADER, {"USD": 1.7}, 0)
    except ValueError as e:
        assert "not found" in str(e)
    else:
        raise AssertionError("a missing template must be refused")


def test_window_pdms_mode_swaps_rows_and_writes_only_the_proposal():
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")
    w = H.generator()
    w.show()
    H.pump(5)
    # normal mode: Contract No visible, PDMS rows hidden
    assert w._contract_no_azn_edit.isVisibleTo(w) and not w._prp_edit.isVisibleTo(w)
    w._pdms_check.setChecked(True)
    assert w._prp_edit.isVisibleTo(w) and not w._contract_no_azn_edit.isVisibleTo(w)
    assert w._attention_edit.text()                     # filled from the request's requester
    # Generate refuses without a PRP No
    out = H.SANDBOX / "out" / "pdms_ui"
    out.mkdir(parents=True, exist_ok=True)
    w._pdms_tpl_edit.setText(str(TEMPLATE))
    w._out_dir_edit.setText(str(out))
    w._job_ref_edit.setText("999")
    H.resolve_disparities(w)
    H.msgs.clear()
    w._generate()
    H.pump(5)
    assert any("PRP No is required" in str(m[-1]) for m in H.msgs), H.msgs
    # with a PRP No: only the .docx (and its PDF, if LibreOffice can) is written
    w._prp_edit.setText("PRP-26-777")
    H.msgs.clear()
    w._generate()
    assert H.wait(lambda: any(m[0] in ("information", "critical") for m in H.msgs), 120)
    assert not any(m[0] == "critical" for m in H.msgs), H.msgs
    assert not list(out.glob("*.xlsx"))
    docs = list(out.glob("PRP-26-777*.docx"))
    assert len(docs) == 1, list(out.iterdir())
    import ctr_tools.cbar_rates as C
    used = [u for u in C.load_usage() if u["prp"] == "PRP-26-777"]       # which document used which rate
    assert len(used) == 1 and used[0]["document"] == docs[0].name and used[0]["rate"] > 0
    d = docx.Document(str(docs[0]))
    assert "{{" not in _text(d) and len(d.tables) >= 2
    w._pdms_check.setChecked(False)
    assert w._contract_no_azn_edit.isVisibleTo(w)       # back to the CTR header


TRANSPORT = [
    {"description": "Minibus+driver", "quantity": 2, "rate_azn": 110.0, "duration": 10, "markup": 0.065, "uom": "Days"},
    {"description": "Fuel for minibus", "quantity": 2, "rate_azn": 10.0, "duration": 10, "markup": 0.0, "uom": "Days"},
    {"description": "Truck (not needed)", "quantity": 0, "rate_azn": 135.0, "duration": 1, "markup": 0.065, "uom": "Trips"},
]


def test_transport_is_listed_after_the_consumables_with_its_cost_only():
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")
    out = B.build_pdms(MAN, [], EQUIP, CONS, TEMPLATE, H.SANDBOX / "pdms_tr", HEADER, {"USD": 1.7}, 0.1,
                       third_party_rows=TRANSPORT)
    d = docx.Document(str(out))
    cons = d.tables[2]
    rows = [[c.text for c in r.cells] for r in cons.rows]
    # header | gloves | Minibus | Fuel | total  (the quantity-0 truck is left out)
    assert len(rows) == 5, rows
    assert rows[1][0] == "GLOVES, SIZE 9"
    assert rows[2] == ["Minibus+driver", "", "", "", f"{2 * 110 * 10 * 1.065:,.2f} ₼"]     # description + cost only
    assert rows[3] == ["Fuel for minibus", "", "", "", "200.00 ₼"]
    consumables = 15 * 2.0 * 1.1 * 1.7
    transport = 2 * 110 * 10 * 1.065 + 200
    assert rows[4][-1] == f"{consumables + transport:,.2f} ₼"                              # counted in the table total
    fills = lambda r: cons.rows[r].cells[0]._tc.tcPr.find(docx.oxml.ns.qn("w:shd")).get(docx.oxml.ns.qn("w:fill"))
    assert fills(2) == fills(1)                                                              # the same look as a consumable row
    grand = 59184 + 11.42 * 62 * 2.5 * 1.7 + consumables + transport
    assert d.tables[3].rows[-1].cells[-1].text == f"{grand:,.2f} ₼"
    assert "{{" not in _text(d)


def test_transport_alone_keeps_the_consumables_table():
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")
    out = B.build_pdms(MAN, [], [], [], TEMPLATE, H.SANDBOX / "pdms_tr2", HEADER, {"USD": 1.7}, 0.1,
                       third_party_rows=TRANSPORT[:1])
    d = docx.Document(str(out))
    assert "Rates are listed in Table № 1 and Table № 2." in _text(d)
    assert [c.text for c in d.tables[1].rows[1].cells][0] == "Minibus+driver"
    try:
        B.build_pdms(MAN, [], [], [], TEMPLATE, H.SANDBOX / "pdms_tr3", HEADER, {"USD": 1.7}, 0.1,
                     third_party_rows=[{"description": "x", "quantity": None, "rate_azn": 1, "duration": 1}])
    except ValueError as e:
        assert "no quantity" in str(e)
    else:
        raise AssertionError("a transport line without a figure must be refused")


def test_hotbed_is_rate_times_days_and_left_out_when_either_is_missing():
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")
    d = docx.Document(str(_build(H.SANDBOX / "pdms_hb1", equip=[], cons=[], hotbed=None)))
    assert "Hotbed" not in _text(d) and "Hotbed" not in "".join(c.text for r in d.tables[0].rows for c in r.cells)
    assert d.tables[0].rows[-1].cells[-1].text == "59,184.00 ₼"                # no hotbed in the total
    d = docx.Document(str(_build(H.SANDBOX / "pdms_hb2", equip=[], cons=[], hotbed={"rate": 70.5, "days": 4})))
    row = [r for r in d.tables[0].rows if r.cells[0].text.startswith("Hotbed")][0]
    assert [c.text for c in row.cells][3] == "4" and row.cells[6].text == "282.00 ₼"
    assert d.tables[0].rows[-1].cells[-1].text == "59,466.00 ₼"


def test_a_manpower_line_without_a_shift_length_is_refused():
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")
    bad = [{**MAN[0], "hours_per_shift": None}]
    try:
        B.build_pdms(bad, [], [], [], TEMPLATE, H.SANDBOX / "pdms_h3", HEADER, {"USD": 1.7}, 0.1)
    except ValueError as e:
        assert "shift length" in str(e)
    else:
        raise AssertionError("days cannot be worked out without the shift length")


def test_each_line_can_have_its_own_currency_converted_before_the_markup_and_rounded_at_the_total():
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")
    rates = {"USD": 1.7, "GBP": 2.25, "EUR": 1.98}
    equip = [{"description": "In GBP", "quantity": 3, "unit": "EA", "rate_per_day": 10.0, "days": 7, "currency": "GBP"},
             {"description": "In USD (default)", "quantity": 1, "unit": "EA", "rate_per_day": 10.0, "days": 1},
             {"description": "In AZN", "quantity": 1, "unit": "EA", "rate_per_day": 10.0, "days": 1, "currency": "AZN"}]
    cons = [{"long_description": "Euro glove", "local_expect_cost": 3.333, "unit_code": "EA", "quantity": 7, "currency": "EUR"}]
    man = [{**MAN[0], "currency": "GBP"}]                        # an hourly rate quoted in pounds
    tr = [{"description": "Minibus", "quantity": 1, "rate_azn": 100.0, "duration": 3, "markup": 0.1, "currency": "EUR"}]
    out = B.build_pdms(man, [], equip, cons, TEMPLATE, H.SANDBOX / "pdms_fx", HEADER, rates, 0.1,
                       third_party_rows=tr, hotbed={"rate": 20.0, "days": 2, "currency": "GBP"})
    d = docx.Document(str(out))
    labour, eq, cn, summary = d.tables
    assert [c.text for c in labour.rows[2].cells][5:7] == [f"{37.14 * 12 * 2.25:,.2f} ₼", f"{720 * 37.14 * 2.25:,.2f} ₼"]
    assert [r.cells[-1].text for r in eq.rows[1:-1]] == [f"{3 * 7 * 10 * 2.25:,.2f} ₼", "17.00 ₼", "10.00 ₼"]
    # consumable: converted first (x1.98), mark-up on the converted price (x1.1), the line total rounded once
    assert cn.rows[1].cells[-1].text == f"{round(7 * 3.333 * 1.98 * 1.1, 2):,.2f} ₼"
    assert cn.rows[2].cells[-1].text == f"{round(1 * 100 * 1.98 * 3 * 1.1, 2):,.2f} ₼"     # transport: converted, then its own mark-up
    row = [r for r in labour.rows if r.cells[0].text.startswith("Hotbed")][0]
    assert row.cells[5].text == "45.00 ₼" and row.cells[6].text == "90.00 ₼"                  # 20 GBP x 2.25, x 2 days
    # a currency without a rate is refused, whatever else is fine
    try:
        B.build_pdms(MAN, [], [{**equip[0], "currency": "TRY"}], [], TEMPLATE, H.SANDBOX / "pdms_fx2", HEADER, rates, 0.1)
    except ValueError as e:
        assert "No CBAR rate for TRY" in str(e)
    else:
        raise AssertionError("TRY has no rate")
    # category defaults can be changed for a whole category
    out = B.build_pdms(MAN, [], equip[:1], [], TEMPLATE, H.SANDBOX / "pdms_fx3", HEADER, rates, 0.1,
                       currencies={"equipment": "EUR"})
    assert docx.Document(str(out)).tables[1].rows[1].cells[-1].text == f"{3 * 7 * 10 * 2.25:,.2f} ₼"   # its own GBP still wins


def test_an_incomplete_hotbed_is_inserted_as_given_without_a_total():
    if not TEMPLATE.is_file():
        raise H.Skip("PDMS_TEMPLATE_short.docx is not in TEST_Files/CTR creator")

    def hotbed_row(hotbed, name):
        d = docx.Document(str(_build(H.SANDBOX / name, equip=[], cons=[], hotbed=hotbed)))
        rows = [r for r in d.tables[0].rows if r.cells[0].text.startswith("Hotbed")]
        return ([c.text for c in rows[0].cells] if rows else None), d.tables[0].rows[-1].cells[-1].text

    # a rate but no days: the line is there with the rate, no days, no total, and the labour total is unchanged
    row, total = hotbed_row({"rate": 59.0, "days": None}, "pdms_hb_a")
    assert row[0] == "Hotbed – if required" and row[3].strip() in ("", "\\xa0") and row[5] == "59.00 ₼" and row[6].strip() in ("", "\\xa0")
    assert total == "59,184.00 ₼"
    # days but no rate: the days only
    row, total = hotbed_row({"rate": None, "days": 5}, "pdms_hb_b")
    assert row[3] == "5" and row[5].strip() in ("", "\\xa0") and row[6].strip() in ("", "\\xa0") and total == "59,184.00 ₼"
    # neither (just ticked): an empty line
    row, total = hotbed_row({"rate": None, "days": None}, "pdms_hb_c")
    assert row is not None and row[0] == "Hotbed – if required" and total == "59,184.00 ₼"
    # both: calculated
    row, total = hotbed_row({"rate": 59.0, "days": 5}, "pdms_hb_d")
    assert row[6] == "295.00 ₼" and total == "59,479.00 ₼"
    # None = left out
    assert hotbed_row(None, "pdms_hb_e")[0] is None

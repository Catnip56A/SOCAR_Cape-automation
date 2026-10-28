"""The "PDMS overview" sheet of the statistics workbook: same figures as the Word proposal, with the arithmetic in cells."""
import shutil
import subprocess

import openpyxl

import harness as H
import ctr_tools.cbar_rates as C
from ctr_tools import builder_pdms as B
from ctr_tools.pdms_overview import SHEET_NAME, add_overview_sheet

MAN = [{"comment": "Day Shift", "num_employees": 1, "description": "PA Level Foreman-National", "quantity": 720,
        "uom": "Hours", "rate_azn": 37.14, "nationality": "National", "hours_per_shift": 12, "encoding": "NS-NAT-OFF-12"},
       {"comment": "Day Shift", "num_employees": 2, "description": "Scaffolder-National", "quantity": 720,
        "uom": "Hours", "rate_azn": 22.53, "nationality": "National", "hours_per_shift": 12, "currency": "GBP",
        "encoding": "MSU-NAT-OFF-12"}]
EQ = [{"description": "Conventional Scaffold", "quantity": 11.42, "unit": "TON", "rate_per_day": 2.5, "days": 62,
       "stock_code": "SCAF-CONV"}]
CO = [{"long_description": "GLOVES 9", "local_expect_cost": 2.59, "unit_code": "EA", "quantity": 15, "product": "PP8201000009"},
      {"long_description": "Euro glove", "local_expect_cost": 3.333, "unit_code": "EA", "quantity": 7, "currency": "EUR",
       "product": "PP8201000010"}]
TR = [{"description": "Minibus+driver", "quantity": 2, "rate_azn": 110.0, "duration": 10, "markup": 0.065, "uom": "Days"}]
HDR = {"prp_number": "PRP-26-056", "date": "2026-09-09", "customer_name": "Example Customer Ltd",
       "attention_name": "Jane Doe", "scope": "P000001 - Test job", "project_type": "Offshore"}
RATES = {"USD": 1.7, "GBP": 2.25, "EUR": 1.98}


def _sheet(hotbed=None, equip=EQ, cons=CO, tr=TR):
    data = B.compute_pdms(MAN, [], equip, cons, HDR, RATES, 0.1, tr, hotbed, None)
    wb = openpyxl.Workbook()
    wb.active.title = "Statistics"
    addr = add_overview_sheet(wb, data, HDR, {"USD": "2026-10-26", "GBP": None, "EUR": "2026-10-26"})
    return wb, data, addr


def _recalculated(wb, name):
    """The workbook with every formula calculated by LibreOffice Calc (None where it is not installed)."""
    if not shutil.which("soffice"):
        return None
    folder = H.SANDBOX / f"ov_{name}"
    folder.mkdir(exist_ok=True)
    src = folder / "in.xlsx"
    wb.save(src)
    r = subprocess.run(["soffice", f"-env:UserInstallation=file://{folder}/lo", "--headless", "--convert-to", "xlsx",
                        "--outdir", str(folder / "out"), str(src)], capture_output=True, text=True, timeout=120)
    out = folder / "out" / "in.xlsx"
    return openpyxl.load_workbook(out, data_only=True) if out.is_file() else None


def test_the_sheet_is_first_and_laid_out_like_the_request_sheet():
    wb, data, addr = _sheet(hotbed={"rate": 59.0, "days": 5, "currency": "AZN"})
    assert wb.sheetnames[0] == SHEET_NAME == "PDMS overview"
    ws = wb[SHEET_NAME]
    texts = {c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)}
    for must in ("Table 1 — Labour rates", "Table 2 — Equipment rates", "Table 3 — Consumables costs",
                 "LABOUR TOTAL:", "MATERIALS TOTAL:", "Summary", "GRAND TOTAL:", "Hotbed – if required",
                 "Normal Day Rate (AZN)", "CBAR rates"):
        assert must in texts, must
    # no orange "provisional" colouring and no footer note
    assert not any(c.fill.start_color.rgb in ("00FFF3E0", "FFFFF3E0") for r in ws.iter_rows() for c in r if c.fill.fill_type)
    assert not any(isinstance(c.value, str) and ("Orange" in c.value or c.value.startswith("Written")) for r in ws.iter_rows() for c in r)
    # the arithmetic is in the cells
    assert ws[addr["labour_total"]].value.startswith("=SUM(")
    assert ws[addr["grand_total"]].value.startswith("=SUM(")
    assert any(isinstance(c.value, str) and "VLOOKUP" in c.value for r in ws.iter_rows() for c in r)
    # the rates block: AZN fixed, the others filled, the source says CBAR day or typed
    rates = {ws.cell(row=r, column=8).value: (ws.cell(row=r, column=9).value, ws.cell(row=r, column=10).value)
             for r in range(5, 9)}
    assert rates["AZN"] == (1, "fixed") and rates["USD"] == (1.7, "CBAR 2026-10-26") and rates["GBP"] == (2.25, "typed by hand")


def test_each_rate_shows_when_it_was_last_checked_or_changed_and_by_what():
    data = B.compute_pdms(MAN, [], EQ, CO, HDR, RATES, 0.1, TR, None, None)
    wb = openpyxl.Workbook()
    add_overview_sheet(wb, data, HDR, {
        "USD": {"day": "2026-10-26", "stamp": "2026-10-27 09:15", "action": "Fetched from CBAR"},
        "GBP": {"day": None, "stamp": "2026-10-27 10:02", "action": "Typed by hand"},
        "EUR": {"day": "2026-10-26", "stamp": "", "action": "Default value, not fetched yet"}})
    ws = wb[SHEET_NAME]
    head = [ws.cell(row=4, column=c).value for c in range(8, 13)]
    assert head == ["Currency", "AZN per 1", "Source", "Last checked / changed", "Last action"]
    rows = {ws.cell(row=r, column=8).value: [ws.cell(row=r, column=c).value for c in (10, 11, 12)] for r in range(5, 9)}
    assert rows["USD"] == ["CBAR 2026-10-26", "2026-10-27 09:15", "Fetched from CBAR"]
    assert rows["GBP"] == ["typed by hand", "2026-10-27 10:02", "Typed by hand"]
    assert rows["EUR"][1:] == ["—", "Default value, not fetched yet"]                 # never checked: a dash, not a made-up date
    assert rows["AZN"] == ["fixed", None, "—"]


def test_stock_codes_and_the_encoding_are_shown_next_to_their_lines():
    wb, data, addr = _sheet(hotbed={"rate": 59.0, "days": 5, "currency": "AZN"})
    ws = wb[SHEET_NAME]

    def table(title):
        r0 = next(r for r in range(1, ws.max_row) if str(ws.cell(row=r, column=2).value or "").startswith(title))
        heads = {ws.cell(row=r0 + 1, column=c).value: c for c in range(2, 14) if ws.cell(row=r0 + 1, column=c).value}
        return r0, heads
    r0, heads = table("Table 1")                                                    # labour: Encoding, right after the description
    assert list(heads)[:2] == ["DESCRIPTION", "Encoding"]
    assert [ws.cell(row=r0 + 2 + i, column=heads["Encoding"]).value for i in range(2)] == ["NS-NAT-OFF-12", "MSU-NAT-OFF-12"]
    hot = next(r for r in range(r0, ws.max_row) if ws.cell(row=r, column=2).value == "Hotbed – if required")
    assert ws.cell(row=hot, column=heads["Encoding"]).value is None               # the Hotbed has none
    r0, heads = table("Table 2")                                                    # equipment: Stock code
    assert list(heads)[:2] == ["Description", "Stock code"]
    assert ws.cell(row=r0 + 2, column=heads["Stock code"]).value == "SCAF-CONV"
    r0, heads = table("Table 3")                                                    # consumables: Stock code
    assert list(heads)[:2] == ["Material Description", "Stock code"]
    assert [ws.cell(row=r0 + 2 + i, column=heads["Stock code"]).value for i in range(2)] == ["PP8201000009", "PP8201000010"]
    # the Word proposal is client-facing and stays as in the sample: no code column there
    d = B.compute_pdms(MAN, [], EQ, CO, HDR, RATES, 0.1, TR, None, None)
    assert d["equipment"][0]["code"] == "SCAF-CONV" and d["consumables"][1]["code"] == "PP8201000010"
    assert d["labour"][0]["encoding"] == "NS-NAT-OFF-12"


def test_the_formulas_give_the_same_totals_as_the_word_proposal():
    wb, data, addr = _sheet(hotbed={"rate": 59.0, "days": 5, "currency": "AZN"})
    calc = _recalculated(wb, "full")
    if calc is None:
        raise H.Skip("LibreOffice Calc is not installed")
    ws = calc[SHEET_NAME]
    assert abs(ws[addr["labour_total"]].value - data["labour_total"]) < 0.005
    assert abs(ws[addr["equipment_total"]].value - data["equip_total"]) < 0.005
    assert abs(ws[addr["materials_total"]].value - data["cons_total"]) < 0.005       # consumables + transport, as in the Word table
    assert abs(ws[addr["grand_total"]].value - data["grand_total"]) < 0.005
    assert data["grand_total"] == round(data["labour_total"] + data["equip_total"] + data["cons_total"], 2)


def test_an_incomplete_hotbed_is_there_without_a_total_and_empty_sections_are_left_out():
    wb, data, addr = _sheet(hotbed={"rate": 59.0, "days": None, "currency": "AZN"}, equip=[], cons=[], tr=[])
    ws = wb[SHEET_NAME]
    texts = {c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)}
    assert "Hotbed – if required" in texts and "Table 2 — Equipment rates" not in texts and "MATERIALS TOTAL:" not in texts
    assert "equipment_total" not in addr and "materials_total" not in addr
    calc = _recalculated(wb, "hb")
    if calc is not None:
        c = calc[SHEET_NAME]
        row = next(r for r in range(1, c.max_row) if c.cell(row=r, column=2).value == "Hotbed – if required")
        assert c.cell(row=row, column=12).value in (None, "")                        # no days: no total
        assert abs(c[addr["labour_total"]].value - data["labour_total"]) < 0.005
        assert abs(c[addr["grand_total"]].value - data["labour_total"]) < 0.005      # only labour is left


def test_changing_a_rate_changes_the_totals_in_the_sheet():
    wb, data, addr = _sheet()
    ws = wb[SHEET_NAME]
    gbp_row = next(r for r in range(5, 9) if ws.cell(row=r, column=8).value == "GBP")
    ws.cell(row=gbp_row, column=9).value = 3.0                                       # the GBP rate, edited in the sheet
    calc = _recalculated(wb, "edit")
    if calc is None:
        raise H.Skip("LibreOffice Calc is not installed")
    expected = round(26740.8 + 2 * 60 * 270.36 * 3.0, 2)
    assert abs(calc[SHEET_NAME][addr["labour_total"]].value - expected) < 0.005


def test_statistics_export_in_pdms_mode_has_the_sheet_first():
    from PySide6.QtWidgets import QMessageBox
    from ctr_tools import window as W
    _fresh = H.SANDBOX / "ov_export"
    _fresh.mkdir(exist_ok=True)
    C._CURRENCIES_PATH = H.SANDBOX / "ov_currencies.json"
    w = H.generator()
    w.show()
    w._pdms_check.setChecked(True)
    w._prp_edit.setText("PRP-26-321")
    H.resolve_disparities(w)

    class _Dialog:                                            # answers the "where to save" question
        def __init__(self, *a, **k): pass
        def exec(self): return W.QDialog.DialogCode.Accepted
        def folder(self): return str(_fresh)
        def deleteLater(self): pass
    real = W._SaveLocationDialog
    W._SaveLocationDialog = _Dialog
    H.msgs.clear()
    try:
        w._save_stats_excel()
    finally:
        W._SaveLocationDialog = real
    assert not any(m[0] == "critical" for m in H.msgs), H.msgs
    files = sorted(_fresh.glob("CTR_Statistics_*.xlsx"))
    assert files, H.msgs
    wb = openpyxl.load_workbook(files[-1])
    assert wb.sheetnames[0] == "PDMS overview" and "Statistics" in wb.sheetnames and "Info" in wb.sheetnames
    assert wb["PDMS overview"]["B1"].value == "PDMS overview — PRP-26-321"
    ov = wb["PDMS overview"]
    assert ov["K4"].value == "Last checked / changed" and ov["L4"].value == "Last action"
    assert all(ov.cell(row=r, column=12).value for r in range(5, 9))                  # every rate says what last happened to it
    # the normal CTR keeps its workbook as before (the file name has the minute in it: same minute = same file)
    w._pdms_check.setChecked(False)
    W._SaveLocationDialog = _Dialog
    try:
        w._save_stats_excel()
    finally:
        W._SaveLocationDialog = real
    names = openpyxl.load_workbook(sorted(_fresh.glob("CTR_Statistics_*.xlsx"))[-1]).sheetnames
    assert "PDMS overview" not in names and names[0] == "Statistics"


def test_days_that_do_not_divide_evenly_still_give_the_same_cents_as_the_word_proposal():
    rows = [{**MAN[0], "quantity": q, "num_employees": 1 + (q % 3)} for q in (100, 101, 137, 250, 333, 41)]
    data = B.compute_pdms(rows, [], [], [], HDR, RATES, 0.1, [], None, None)
    wb = openpyxl.Workbook()
    addr = add_overview_sheet(wb, data, HDR, {})
    calc = _recalculated(wb, "uneven")
    if calc is None:
        raise H.Skip("LibreOffice Calc is not installed")
    ws = calc[SHEET_NAME]
    sheet_totals = [ws.cell(row=r, column=12).value for r in range(1, ws.max_row)
                    if ws.cell(row=r, column=2).value and "Foreman" in str(ws.cell(row=r, column=2).value)]
    assert [round(v, 2) for v in sheet_totals] == [ln["total"] for ln in data["labour"]]      # line by line, to the cent
    assert abs(ws[addr["labour_total"]].value - data["labour_total"]) < 0.005


def test_a_failing_overview_sheet_does_not_lose_the_normal_statistics_sheets():
    from ctr_tools import window as W
    out = H.SANDBOX / "ov_fail"
    out.mkdir(exist_ok=True)
    C._CURRENCIES_PATH = H.SANDBOX / "ovf_currencies.json"
    w = H.generator()
    w.show()
    w._pdms_check.setChecked(True)
    w._prp_edit.setText("PRP-26-322")
    H.resolve_disparities(w)
    eq = w._req_equip_tbl
    r = next(i for i in range(eq.rowCount()) if eq.item(i, W._EQ_STATUS).text().startswith("✓"))
    for col, val in ((W._EQ_QTY, "2"), (W._EQ_RATE, "10"), (W._EQ_DAYS, "3")):
        eq.item(r, col).setText(val)
    eq.item(r, W._EQ_RATE).setData(W._CUR_ROLE, "GBP")
    w._fx["GBP"]["edit"].setText("")                                                  # a currency in use without a rate

    class _Dialog:
        def __init__(self, *a, **k): pass
        def exec(self): return W.QDialog.DialogCode.Accepted
        def folder(self): return str(out)
        def deleteLater(self): pass
    real = W._SaveLocationDialog
    W._SaveLocationDialog = _Dialog
    H.msgs.clear()
    try:
        w._save_stats_excel()
    finally:
        W._SaveLocationDialog = real
    assert not any(m[0] == "critical" for m in H.msgs), H.msgs                         # no "Save error"
    saved = [m for m in H.msgs if m[0] == "information" and "Statistics saved" in str(m[-1])]
    assert saved and "NOT added" in str(saved[-1][-1]) and "No CBAR rate for GBP" in str(saved[-1][-1])
    names = openpyxl.load_workbook(sorted(out.glob("CTR_Statistics_*.xlsx"))[-1]).sheetnames
    assert "PDMS overview" not in names and names[0] == "Statistics" and "Info" in names   # the rest is there

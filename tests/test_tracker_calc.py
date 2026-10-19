"""MA143: the CTR Tracker's fast write path asks Excel to recalculate on open (fullCalcOnLoad), because it
writes formula text but cannot compute results. Works on temp copies only; the test tracker is read-only."""
import shutil
import zipfile
from datetime import date

from lxml import etree

import harness as H
from ctr_tools import tracker_fast as TF

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _wb(*children):
    body = "".join(children)
    return f'<workbook xmlns="{NS}">{body}</workbook>'.encode()


def _children(xml: bytes):
    return [etree.QName(c).localname for c in etree.fromstring(xml)]


def test_an_existing_calcpr_gets_the_flag_and_keeps_its_other_attributes():
    out = TF._force_full_calc_on_load(_wb("<sheets/>", '<calcPr calcId="191029"/>'))
    calc = etree.fromstring(out).find(f"{{{NS}}}calcPr")
    assert calc.get("fullCalcOnLoad") == "1" and calc.get("calcId") == "191029"
    assert _children(out) == ["sheets", "calcPr"]


def test_a_missing_calcpr_is_created_in_the_right_place():
    out = TF._force_full_calc_on_load(_wb("<fileVersion/>", "<sheets/>", "<definedNames/>", "<extLst/>"))
    assert _children(out) == ["fileVersion", "sheets", "definedNames", "calcPr", "extLst"]
    assert etree.fromstring(out).find(f"{{{NS}}}calcPr").get("fullCalcOnLoad") == "1"
    bare = TF._force_full_calc_on_load(_wb())               # nothing else in the file
    assert _children(bare) == ["calcPr"]


def test_writing_an_entry_marks_the_saved_tracker_for_recalculation_and_it_still_opens():
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci
    from ctr_tools.config import CFG
    from ctr_tools.tracker import CTREntry
    src = H.ROOT / "TEST_Files" / "CTR tracker" / "CTR-Tracker copy.xlsm"
    if not src.exists():
        raise H.Skip("TEST_Files/CTR tracker is not here")
    cfg = CFG["ctr_tracker"]
    work = H.SANDBOX / "tracker_copy_calc.xlsm"
    shutil.copy2(src, work)

    def calc_attrs():
        root = etree.fromstring(zipfile.ZipFile(work).read("xl/workbook.xml"))
        return dict(root.find(f"{{{NS}}}calcPr").attrib)

    assert "fullCalcOnLoad" not in calc_attrs()
    ws = openpyxl.load_workbook(work, data_only=True, keep_vba=True)[cfg["sheet_name"]]
    spare = next((r for r in range(2, ws.max_row + 1)
                  if ws.cell(r, ci(cfg["col_ctr_number"])).value and not ws.cell(r, ci(cfg["col_description"])).value), None)
    if spare is None:
        raise H.Skip("no spare row in the test tracker")
    number = str(ws.cell(spare, ci(cfg["col_ctr_number"])).value)
    entry = CTREntry(client="Test client", ctr_number=number, ctr_date=date(2026, 10, 15),
                     description="calc flag test", value=123.0, currency="USD",
                     tracker_location="Onshore", job_type="x")
    written, skipped, _bak, _notes = TF.write_entries_fast(work, [entry], make_backup=False)
    assert written and not skipped, (written, skipped)
    assert calc_attrs().get("fullCalcOnLoad") == "1"
    assert openpyxl.load_workbook(work, keep_vba=True).sheetnames[0] == cfg["sheet_name"]


def test_nothing_is_written_to_the_file_when_no_entry_could_be_placed():
    src = H.ROOT / "TEST_Files" / "CTR tracker" / "CTR-Tracker copy.xlsm"
    if not src.exists():
        raise H.Skip("TEST_Files/CTR tracker is not here")
    from ctr_tools.tracker import CTREntry
    work = H.SANDBOX / "tracker_copy_untouched.xlsm"
    shutil.copy2(src, work)
    before = work.read_bytes()
    written, skipped, _b, _n = TF.write_entries_fast(
        work, [CTREntry(ctr_number="", currency="USD")], make_backup=False)
    assert not written and skipped
    assert work.read_bytes() == before                      # the file is only touched when something was written

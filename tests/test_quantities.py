"""Quantities of 0 or less: Generate warns; transport lines are hidden in the AZN CTR."""
import openpyxl

import harness as H
from harness import W


def _third_party(out):
    ws = openpyxl.load_workbook(H.output_file(out, "AZN")).worksheets[0]
    rows = {}
    for r in range(1, ws.max_row + 1):
        d = ws.cell(r, 3).value
        if d in ("Minibus+driver", "Fuel for minibus", "Truck"):
            rows[d] = (r, ws.row_dimensions[r].hidden, ws.cell(r, 7).value, ws.cell(r, 2).value)
    titles = {str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1)}
    return rows, titles


def test_transport_zero_and_negative_quantities_are_hidden_with_total_zero():
    w = H.generator()
    tr = w._transport_tbl
    by_desc = {tr.item(r, W._TR_DESC).text(): r for r in range(tr.rowCount())}
    tr.item(by_desc["Fuel for minibus"], W._TR_QTY).setText("0")
    tr.item(by_desc["Truck"], W._TR_QTY).setText("-1")
    out, msgs = H.generate(w, "tr_zero")
    text = H.all_text(msgs)
    assert "HIDDEN" in text and "Fuel for minibus" in text and "Truck" in text
    rows, _ = _third_party(out)
    assert not rows["Minibus+driver"][1], "a normal line must stay visible"
    for d in ("Fuel for minibus", "Truck"):
        assert rows[d][1] is True, f"{d} should be hidden"
        assert rows[d][2] == 0, f"{d} total must be 0 so a hidden line never changes the total"


def test_all_transport_lines_zero_leaves_the_section_out():
    w = H.generator()
    tr = w._transport_tbl
    for r in range(tr.rowCount()):
        tr.item(r, W._TR_QTY).setText("0")
    out, _ = H.generate(w, "tr_all_zero")
    rows, titles = _third_party(out)
    assert not rows
    assert "Third Party Activities" not in titles


def test_other_nonpositive_quantities_only_warn():
    w = H.generator()
    w._req_equip_tbl.item(0, W._EQ_QTY).setText("0")
    w._req_manpower_tbl.item(0, W._MP_NUMEMP).setText("0")
    out, msgs = H.generate(w, "qty_zero")
    text = H.all_text(msgs)
    assert "0 or less" in text and "Num Employees 0" in text
    # the rows are still in the CTR, as entered
    desc = w._req_manpower_tbl.item(0, W._MP_DESC).text()
    ws = openpyxl.load_workbook(H.output_file(out, "AZN")).worksheets[0]
    assert any(ws.cell(r, 3).value == desc for r in range(1, ws.max_row + 1))

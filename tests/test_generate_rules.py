"""Generate: nothing is assumed, incomplete rows are left out, builders refuse to guess."""
import openpyxl

import cost_stats
import harness as H
from harness import W


def _descs(path, sheet_index=0, col=3):
    ws = openpyxl.load_workbook(path).worksheets[sheet_index]
    return {str(r[col - 1].value) for r in ws.iter_rows() if r[col - 1].value}


def _pricing(out):
    return openpyxl.load_workbook(H.output_file(out, "USD"))["Pricing"]


def test_incomplete_scaffold_is_listed_and_left_out():
    w = H.generator()
    cells = w._incomplete_cells()
    assert any("Scaffold" in c["desc"] for c in cells), [c["desc"] for c in cells]
    out, msgs = H.generate(w, "scaffold")
    assert "NOT be included" in H.all_text(msgs)
    texts = {str(r[1].value) for r in _pricing(out).iter_rows() if r[1].value}
    assert not any("Conventional Scaffold" in t for t in texts)


def test_blank_values_leave_exactly_those_rows_out():
    w = H.generator()
    mp, cs, tr = w._req_manpower_tbl, w._req_cons_tbl, w._transport_tbl
    mp_desc = mp.item(0, W._MP_DESC).text()
    cs_desc = cs.item(0, W._CS_DESC).text()
    tr_desc = tr.item(0, W._TR_DESC).text()
    mp.item(0, W._MP_NUMEMP).setText("")
    cs.item(0, W._CS_PRICE).setText("")
    tr.item(0, W._TR_QTY).setText("")
    out, msgs = H.generate(w, "blank")
    text = H.all_text(msgs)
    for part in (mp_desc, cs_desc, tr_desc):
        assert part in text, f"{part!r} not named in the warning"
    assert mp_desc not in _descs(H.output_file(out, "AZN"))
    assert tr_desc not in _descs(H.output_file(out, "AZN"))
    assert cs_desc not in {str(r[1].value) for r in _pricing(out).iter_rows() if r[1].value}


def test_nonrecharg_needs_no_rate_but_blank_rate_does():
    w = H.generator()
    eq = w._req_equip_tbl
    r = next(i for i in range(eq.rowCount()) if eq.item(i, W._EQ_STATUS).text().startswith("✓")
             and not eq.item(i, W._EQ_DESC).data(W._SCAFFOLD_ROLE))
    eq.item(r, W._EQ_RATE).setText("NONRECHARG")
    assert not [c for c in w._incomplete_cells() if c["tbl"] is eq and c["row"] == r and c["field"] == "Rate"]
    eq.item(r, W._EQ_RATE).setText("")
    assert [c for c in w._incomplete_cells() if c["tbl"] is eq and c["row"] == r and c["field"] == "Rate"]


def test_builders_refuse_to_guess():
    from ctr_tools import builder_azn as Z, builder_usd as U
    for fn, row, key in ((U._require_float, {"quantity": ""}, "quantity"),
                         (U._require_float, {"quantity": None}, "quantity"),
                         (U._require_rate, {"rate": "0x"}, "rate"),
                         (Z._require_float, {"quantity": "abc"}, "quantity")):
        try:
            fn(row, key, "line")
        except ValueError:
            continue
        raise AssertionError(f"{fn.__name__} accepted {row}")
    assert U._require_rate({"r": "nonrecharg"}, "r", "x") == "NONRECHARG"
    assert U._require_float({"q": 0}, "q", "x") == 0.0       # a stated 0 is a value
    try:
        U._require_text({"unit": " "}, "unit", "x")
    except ValueError:
        pass
    else:
        raise AssertionError("blank unit accepted")


def test_stock_code_not_in_sage_gets_an_empty_cost_cell_until_typed():
    w = H.generator()
    target = None
    for tbl, c_desc, c_match, c_code in ((w._req_equip_tbl, W._EQ_DESC, W._EQ_MATCH, W._EQ_CODE),
                                         (w._req_cons_tbl, W._CS_DESC, W._CS_MATCH, W._CS_CODE)):
        for r in range(tbl.rowCount()):
            d = tbl.item(r, c_desc)
            if d.data(W._SAGE_COST_ROLE) is None and not d.data(W._SCAFFOLD_ROLE) \
                    and tbl.item(r, W._EQ_STATUS if tbl is w._req_equip_tbl else W._CS_STATUS).text().startswith("✓"):
                key = tbl.item(r, c_match).text().strip() or tbl.item(r, c_code).text().strip() or d.text().strip()
                target = (d.text().strip(), key)
    assert target, "the test request has no Stock Code that SAGE lacks"
    desc, key = target

    def sc_cell(out):
        for row in _pricing(out).iter_rows():
            if row[1].value and str(row[1].value) == desc:
                return row[8].value
        raise AssertionError(f"{desc!r} not in the Pricing sheet")

    out, msgs = H.generate(w, "nosage")
    assert "NOT IN SAGE" in H.all_text(msgs)
    assert sc_cell(out) is None
    w._stats_panel._manual[cost_stats.norm_key(key)] = 12.5
    out, _ = H.generate(w, "nosage_typed")
    assert str(sc_cell(out)).endswith("*12.5"), sc_cell(out)

"""
ctr_tools/builder_azn.py

Copies the AZN CTR template and writes labor rows into sheet "Main AZN".

Row/column positions are read from ctr_tools/template_config.json
(section "azn_template") so they can be adjusted without touching this
file when a new template version shifts the layout.

Default layout (217_AZN template):
  Row 3  : E3 = CTR ref string, O3 = job ref integer
  Row 7  : "Project Support" section header (fixed text)
  Row 8  : column headers for the Project Support section
  Rows 9–18 : "support" section data rows — any manpower row whose
    description contains the keyword "support" or names one of the
    always-support roles (case-insensitive; see is_support_manpower),
    regardless of its Onshore/Offshore type
  Row 20 : G20 = total project support
  Row 21 : section header for every other manpower row — text is
    "Onshore Activities" or "Offshore Activities" depending on the CTR
    request's project type (see activities_label), not fixed template text
  Row 22 : column headers for that section
  Rows 23–62: data rows for every manpower row NOT matched as "support"
  Rows 65–68: Comments box (merged A65:G68) — the CTR Request's required
    Additional Information items (Floatel, Accommodation, Per Diem, …),
    written here by _write_comments_extra; left as the template's plain
    "Comments:" placeholder when nothing is required. See activities_label
    for what used to happen instead (appended onto the row-21 header).
  Row 69 : G69 = total for the Onshore/Offshore Activities section
  Optional "Third Party Activities" section, inserted right after that
    total when the CTR Request has transport / hired-service lines — the
    template ships without it (see _write_third_party_section), and the
    Summary block below gains a matching "Total Other Activities" line.
  Summary (rows 70–76, before any of the above shifts them down):
    G72 = total project support
    G73 = total of the other section
    G76 = estimated CTR total

  Editable header fields (rows 3–6, shared layout with USD template):
    B3 = Client, C3 = Sub-Client, B4 = Location, E4 = Date,
    G4 = Contract No, E5 = Revision, A6 = Scope / description (merged A6:G6),
    C1 = Comments, carried over from the CTR Request's own free-text
    Comments box — unrelated to the A65:G68 Comments box above, which
    holds only the required Additional Information items

Input values (comment, employees, quantity, rate, etc.) are written as
plain Python values. Anything derived from another cell — row totals,
section totals, summary cells — is written as an Excel formula referencing
that cell, so editing an input in Excel recalculates its dependents. This
only matters for the .xlsx output; the PDF export is a flat rendering, so
formulas there simply show their last-calculated value.

The "support" keyword itself (see is_support_manpower / strip_support_keyword)
is an internal routing marker, not client-facing text — it's stripped from
the Comment and Description cells of every written row before they reach
the document, so "Painter Support" appears as plain "Painter". A row routed
there by its role name instead keeps its description as written, since
there's no marker word to remove.
"""

from __future__ import annotations

import re
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string as _col_idx
from openpyxl.utils import get_column_letter as _col_letter
from openpyxl.worksheet.properties import PageSetupProperties

from ctr_tools.config import CFG
from ctr_tools.naming import ctr_output_filename

_azn = CFG["azn_template"]

_SUPPORT_DATA_START = _azn["onshore_data_start"]
_SUPPORT_DATA_END   = _azn["onshore_data_end"]
_SUPPORT_TOTAL_GAP  = _azn["onshore_total_gap"]   # blank rows between data end and total

_OTHER_DATA_START = _azn["offshore_data_start"]
_OTHER_DATA_END   = _azn["offshore_data_end"]
_OTHER_TOTAL_GAP  = _azn["offshore_total_gap"]

_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y")

# Manpower rows whose description contains this keyword (any casing —
# "Support", "SUPPORT", "Supporting", "Life Support", etc. all match as a
# plain substring) go into the fixed "Project Support" rows (9–18); every
# other manpower row goes into the section starting at row 23.
_SUPPORT_KEYWORD = "support"
_SUPPORT_RE = re.compile(re.escape(_SUPPORT_KEYWORD), re.IGNORECASE)

# Roles that belong in "Project Support" whatever the request calls them —
# a project engineer is charged there whether or not the requester thought
# to write "support" next to the role. Matched as a plain lowercase
# substring, same as the keyword above, so "Senior Project Engineer" and
# "PROJECT ENGINEER-National" both match one "project engineer" entry.
# Keep entries to full role names (azn_template.support_roles in
# template_config.json): a short fragment would match inside unrelated
# roles and silently move them out of the activities section.
_SUPPORT_ROLES = tuple(
    role.strip().lower()
    for role in _azn.get("support_roles", [])
    if str(role).strip()
)


def is_support_manpower(description: str) -> bool:
    """
    True if a manpower description should be treated as "support" manpower —
    either because it carries the "support" keyword, or because it names one
    of the always-support roles in _SUPPORT_ROLES. Any one of those matching
    is enough.
    """
    desc = (description or "").lower()
    return _SUPPORT_KEYWORD in desc or any(role in desc for role in _SUPPORT_ROLES)


def strip_support_keyword(text: str) -> str:
    """
    Removes the "support" keyword from a manpower match string before it's
    looked up in the pricebook — a role like "Scaffolder Support" should
    match the base pricebook item "Scaffolder" rather than fail as an
    unmatched exact string. Cleans up whitespace/punctuation left behind by
    the removal (e.g. "Painter (Support)" → "Painter"); falls back to the
    original text if stripping would leave nothing.
    """
    if not text:
        return text
    stripped = _SUPPORT_RE.sub("", text)
    stripped = re.sub(r"\(\s*\)", "", stripped)
    stripped = re.sub(r"\s{2,}", " ", stripped).strip(" -_/,()")
    return stripped or text.strip()


def activities_label(project_type: str) -> str:
    """
    Section header text for the non-"support" manpower section (row 21/H21):
    "Onshore Activities" or "Offshore Activities", chosen from the CTR
    request's project type. Falls back to "Offshore Activities" (the
    template's original hardcoded text) when project_type is blank or
    doesn't recognize either word.

    The CTR Request's required Additional Information items (Floatel,
    Accommodation, Per Diem, …) used to be appended to this same header —
    see _write_comments_extra for where that now lives instead, in its own
    dedicated Comments box (A65:G68) rather than folded into this label.
    """
    if "onshore" in (project_type or "").lower():
        return _azn.get("label_onshore_activities", "Onshore Activities")
    return _azn.get("label_offshore_activities", "Offshore Activities")


def _write_comments_extra(ws, addr: str, row_shift: int, required_info: list[str] | None) -> None:
    """
    Writes the CTR Request's required Additional Information items (Floatel,
    Accommodation, Per Diem, …) into the dedicated Comments box at `addr`
    (A65:G68 by default) — see activities_label, which used to fold these
    into the row-21 section header instead.

    `addr` is a merged cell, so only its top-left anchor actually holds a
    value; `row_shift` accounts for any support/offshore overflow rows
    inserted above it (the merge itself shifts down with them). Left
    untouched — keeping the template's plain "Comments:" placeholder —
    when there's nothing required, the same way an optional header field
    elsewhere in this module leaves the template's own value alone.
    """
    items = [str(item).strip() for item in (required_info or []) if str(item).strip()]
    if not items:
        return
    row, col = _cell_row_col(addr)
    cell = _anchor_cell(ws, row + row_shift, col)
    if cell is not None:
        cell.value = f"Comments:\n{', '.join(items)} required"


def _cell_row_col(addr: str) -> tuple[int, int]:
    """Convert a cell address like 'G67' to (row=67, col=7)."""
    col_str = "".join(c for c in addr if c.isalpha())
    row_str = "".join(c for c in addr if c.isdigit())
    return int(row_str), _col_idx(col_str)


def _anchor_cell(ws, row: int, col: int):
    """
    The cell that actually stores what's displayed at (row, col).

    openpyxl represents every cell of a merged range except its top-left as
    a read-only MergedCell — assigning to one raises "'MergedCell' object
    attribute 'value' is read-only". Templates are re-merged by hand
    between revisions, so any cell this module addresses by fixed position
    can quietly end up inside somebody else's merge. Returns the range's
    anchor in that case (which is where the value visibly lands anyway), or
    None if the merge can't be resolved.
    """
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        return cell
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(row=rng.min_row, column=rng.min_col)
    return None


def _write_addr(ws, addr: str, value) -> None:
    """Write `value` at `addr`, following a merge to the cell that holds it."""
    cell = _anchor_cell(ws, *_cell_row_col(addr))
    if cell is not None:
        cell.value = value


def _set_cell(ws, row: int, col: int, value) -> None:
    """Write value at (row, col), skipping non-top-left merged-cell slots.

    Used for the bulk section writes, where the anchor of any merge inside
    the block is visited by the loop in its own right — so skipping is both
    safe and avoids clearing a merge that reaches outside the block."""
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        cell.value = value


def _insert_rows_preserving_merges(ws, insert_row: int, amount: int) -> None:
    """
    openpyxl's insert_rows() shifts cell values down but leaves merged-cell
    ranges at their original row numbers — a stale merge left behind at the
    insertion point (e.g. a section-header row) silently swallows any data
    later written into whichever row now occupies that position. Unmerging
    *after* insert_rows() also fails (the merge's cached cell objects no
    longer match reality), so affected ranges must be unmerged first,
    inserted around, then re-merged at their shifted position.
    """
    affected = [
        (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
        for mr in ws.merged_cells.ranges if mr.min_row >= insert_row
    ]
    for min_row, min_col, max_row, max_col in affected:
        ws.unmerge_cells(start_row=min_row, start_column=min_col,
                          end_row=max_row, end_column=max_col)

    ws.insert_rows(insert_row, amount)

    for min_row, min_col, max_row, max_col in affected:
        ws.merge_cells(start_row=min_row + amount, start_column=min_col,
                        end_row=max_row + amount, end_column=max_col)


def _copy_row_format(ws, src_row: int, dst_row: int, col_count: int) -> None:
    """Copy cell formatting (borders, fill, font, alignment) from one row to another."""
    src_dim = ws.row_dimensions.get(src_row)
    if src_dim and src_dim.height:
        ws.row_dimensions[dst_row].height = src_dim.height
    for col in range(1, col_count + 1):
        src = ws.cell(row=src_row, column=col)
        dst = ws.cell(row=dst_row, column=col)
        if src.has_style:
            dst.font         = copy(src.font)
            dst.border       = copy(src.border)
            dst.fill         = copy(src.fill)
            dst.number_format = src.number_format
            dst.alignment    = copy(src.alignment)


def _clear_placeholder_cells(ws, texts) -> int:
    """
    Blank every cell whose entire value is one of `texts`.

    The templates carry filler strings in their helper cells — "AA" all the
    way down the column-H mirror on the Main sheets, "aa" beside the
    Pricing sheet's margin block. They mean nothing; they normally go
    unnoticed because column H is hidden and those Pricing rows fall
    outside the print area, so they reappear as soon as a template is saved
    with the column shown. Since a generated CTR goes to a client, they're
    stripped rather than passed on.

    Matching is on the whole trimmed value, so the genuine content sharing
    those columns — the "=A10"/"=IF(B8=...)" mirror formulas, the real
    labels — is never touched. Returns how many cells were cleared.
    """
    wanted = {str(t).strip() for t in (texts or []) if str(t).strip()}
    if not wanted:
        return 0
    cleared = 0
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() in wanted:
                cell.value = None
                cleared += 1
    return cleared


def _extend_print_area(ws, extra_rows: int) -> None:
    """
    Push the sheet's print area down by `extra_rows`.

    The print area is a fixed range saved in the template ("$A$1:$G$77")
    and openpyxl does not move it when rows are inserted — so every row a
    generated CTR adds pushes that much of the document out the bottom of
    what actually prints. It is silent in Excel and only shows up in the
    PDF, where the tail of the document (the Estimated CTR Total line and
    the signature block) simply isn't there.

    Left alone if the sheet has no print area, or has a multi-range one
    this can't safely reason about — better to print the template's range
    than to guess wrong about a layout that isn't the one this understands.
    """
    area = ws.print_area
    if not area or extra_rows <= 0:
        return
    if isinstance(area, (list, tuple)):
        if len(area) != 1:
            return
        area = area[0]
    if "," in area or ":" not in area:
        return

    start, end = area.rsplit(":", 1)
    digits, cut = "", len(end)
    while cut > 0 and end[cut - 1].isdigit():
        cut -= 1
        digits = end[cut] + digits
    if not digits:
        return
    ws.print_area = f"{start}:{end[:cut]}{int(digits) + extra_rows}"


def _fit_to_page_width(ws) -> None:
    """
    Forces the sheet to print at exactly one page wide (any number of pages
    tall), overriding the template's fixed print scale — a fixed scale
    percentage can fit every column on one page width when rendered by one
    Excel/LibreOffice install but overflow onto a second page in another,
    splitting every row's right-hand columns off onto an unreadable
    separate page. "Fit to 1 page wide" is computed by the renderer at
    export time instead, so it's correct regardless of font metrics.
    """
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if ws.sheet_properties.pageSetUpPr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties()
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # The template's right margin is 0" (left barely 0.15-0.24") — with
    # content already forced to fill exactly one page's width, that leaves
    # no breathing room at all on the sides. A small fixed margin here
    # still fits on one page (fitToWidth shrinks to whatever's left inside
    # the margins) while giving the page visible whitespace on both edges.
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3


def _safe_float(value, default: float = 0.0) -> float:
    """Best-effort float conversion; `default` for blanks, non-numbers, and
    NaN (a NaN here would otherwise format as the literal text "nan" inside
    an Excel formula string — e.g. the Third Party section's
    qty*rate*duration*(1+markup) — producing a broken formula on open)."""
    try:
        f = float(value)
        return f if f == f else default   # NaN guard
    except (TypeError, ValueError):
        return default


def _parse_date(value):
    """Best-effort string → datetime conversion; returns value unchanged if unparseable."""
    if isinstance(value, str):
        for fmt in _DATE_FORMATS:
            try:
                return datetime.strptime(value, fmt)
            except ValueError:
                continue
    return value


# Row-height estimate for wrapped comment text. Deliberately rough and
# deliberately pessimistic: a column's width is expressed in characters of
# the workbook's *default* font, so measuring a smaller cell font against it
# under-counts how many characters fit on a line, over-counts the lines, and
# leaves the row slightly taller than strictly needed. A little too tall
# wastes a few points of page; too short silently cuts the comment off.
_DEFAULT_ROW_HEIGHT     = 15.0
_DEFAULT_FONT_SIZE      = 11.0
_LINE_HEIGHT_FACTOR     = 1.35    # point size -> line box height
_MAX_COMMENT_ROW_HEIGHT = 200.0


def _wrapped_row_height(text: str, col_width, font_size, base_height: float) -> float:
    """Height in points for `text` wrapped inside a column `col_width` wide."""
    chars_per_line = max(10, int(col_width or _DEFAULT_ROW_HEIGHT))
    lines = sum(
        max(1, -(-len(para) // chars_per_line))          # ceil division
        for para in str(text).split("\n")
    )
    needed = max(1, lines) * (font_size or _DEFAULT_FONT_SIZE) * _LINE_HEIGHT_FACTOR + 4
    return max(base_height, min(needed, _MAX_COMMENT_ROW_HEIGHT))


def _write_comments(ws, addr: str, text: str) -> None:
    """
    Write the CTR Request's comments into `addr`, forcing the cell to read
    horizontally and wrap.

    Both templates style that cell with textRotation=90 and no wrap — it
    sits in the tall banner row beside the standing PO note, where it was
    only ever meant to hold a short sideways tag. Left as-is, a real
    sentence of comments comes out running vertically down the page and
    clipped at the column edge. The row is also grown, if needed,
    to fit however many lines the wrapped text takes — never shrunk, so the
    PO note sharing row 1 keeps its space, and never touched at all when
    there are no comments.
    """
    # The comments cell is addressed by fixed position too, so follow a
    # merge to whatever cell actually holds it (see _anchor_cell).
    cell = _anchor_cell(ws, *_cell_row_col(addr))
    if cell is None:
        return
    cell.value = text
    cell.alignment = Alignment(
        horizontal="left", vertical="center", wrap_text=True, textRotation=0)
    if not text:
        return

    # Give the row an explicit height tall enough for the wrapped text,
    # never shorter than whatever the template set (row 1 also carries the
    # PO note, which needs its own space). Clearing the height so the row
    # auto-fits instead does NOT work here: Excel would auto-fit on open,
    # but LibreOffice — which renders the PDF export — just draws the row at
    # its default height, and the whole banner row came out as a sliver.
    row = cell.row
    col = cell.column_letter
    base  = ws.row_dimensions[row].height or _DEFAULT_ROW_HEIGHT
    width = ws.column_dimensions[col].width if col in ws.column_dimensions else None
    ws.row_dimensions[row].height = _wrapped_row_height(
        text, width, cell.font.size, base)


def _write_third_party_section(
    ws, after_row: int, rows: list[dict], col_count: int, labor_data_start: int,
) -> tuple[int, int]:
    """
    Insert a "Third Party Activities" section immediately below `after_row`
    (the activities section's Total row) and write `rows` into it.

    The shipped AZN template has no such section — a CTR that needs one
    (transport, hired services) has it added by hand, between the labor
    total and the Summary block, laid out as:

        Third Party Activities                          <- merged A:G
        Comment | Quantity | Description | Mark Up, % | Duration&UOM | Rate | Total
        ...one row per line item...
        Total Third Party Activities            =SUM() <- merged A:F

    Returns (total_row, rows_inserted). Formatting for each new row is
    copied from the equivalent row of the labor section starting at
    `labor_data_start` — its title, column-header, data and total rows are
    all above the insertion point, so they're still where the caller found
    them (which is not where the template had them, if the request needed
    extra manpower rows inserted earlier).

    Each row's total is qty x rate x duration x (1 + mark-up). Duration is
    embedded in the formula as a literal because its own cell is text ("6
    Days" / "2 Trips") — the request states a duration but not what a
    duration means for that vehicle, so it stays human-readable rather than
    being split into a number and a unit column the template doesn't have.
    """
    hdr_row     = after_row + 1
    col_hdr_row = hdr_row + 1
    data_start  = col_hdr_row + 1
    data_end    = data_start + len(rows) - 1
    total_row   = data_end + 1
    inserted    = len(rows) + 3

    _insert_rows_preserving_merges(ws, hdr_row, inserted)

    # Source rows for formatting: the labor section's own title / column
    # header / data / total rows, all above the insertion point.
    src_total   = after_row
    src_col_hdr = labor_data_start - 1
    src_data    = labor_data_start
    src_hdr     = labor_data_start - 2

    _copy_row_format(ws, src_hdr,     hdr_row,     col_count)
    _copy_row_format(ws, src_col_hdr, col_hdr_row, col_count)
    for r in range(data_start, data_end + 1):
        _copy_row_format(ws, src_data, r, col_count)
    _copy_row_format(ws, src_total, total_row, col_count)

    ws.merge_cells(start_row=hdr_row,   start_column=1, end_row=hdr_row,   end_column=7)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)

    ws.cell(row=hdr_row, column=1).value = _azn.get(
        "third_party_label", "Third Party Activities")
    for i, heading in enumerate(_azn.get("third_party_headers", []), start=1):
        ws.cell(row=col_hdr_row, column=i).value = heading

    for i, tp in enumerate(rows):
        r = data_start + i
        qty      = _safe_float(tp.get("quantity", 1), 1.0)
        rate     = _safe_float(tp.get("rate_azn", 0), 0.0)
        markup   = _safe_float(tp.get("markup", 0), 0.0)
        duration = _safe_float(tp.get("duration", 1), 1.0)
        uom      = str(tp.get("uom", "") or "").strip()

        ws.cell(row=r, column=1).value = str(tp.get("comment", "") or "")
        ws.cell(row=r, column=2).value = qty
        ws.cell(row=r, column=3).value = str(tp.get("description", "") or "")
        # A line that carries no mark-up is left blank rather than showing
        # "0.0%", the way the hand-priced CTR this section is modelled on
        # writes its un-marked-up fuel line. Excel reads the empty cell as 0
        # in the row's (1+D) term either way.
        markup_cell = ws.cell(row=r, column=4)
        markup_cell.value = markup or None
        markup_cell.number_format = _azn.get("third_party_markup_format", "0.0%")
        ws.cell(row=r, column=5).value = f"{duration:g} {uom}".strip()
        ws.cell(row=r, column=6).value = rate
        ws.cell(row=r, column=7).value = f"=B{r}*F{r}*{duration:g}*(1+D{r})"

    for r in range(hdr_row, total_row + 1):
        ws.row_dimensions[r].hidden = False

    ws.cell(row=total_row, column=1).value = _azn.get(
        "third_party_total_label", "Total Third Party Activities")
    ws.cell(row=total_row, column=7).value = f"=SUM(G{data_start}:G{data_end})"

    return total_row, inserted


def build_azn(
    support_rows: list[dict],
    other_rows: list[dict],
    template_path: str | Path,
    output_dir: str | Path,
    job_ref: str,
    header: dict | None = None,
    third_party_rows: list[dict] | None = None,
) -> Path:
    """
    Write labor rows into a copy of the AZN template.

    support_rows: manpower rows matched as support manpower — by the
        "support" keyword or by role name (see is_support_manpower) —
        written into rows 9–18.
    other_rows: every other manpower row — written into rows 23+.

    Each row is a dict with keys:
        comment (str — the manpower table's Shift/Shift Type, e.g.
            "Day Shift, Normal"; see window.py's Generate handler),
        num_employees (int/float), description (str),
        quantity (float), uom (str), rate_azn (float), nationality (str)

    header: optional dict with keys client, sub_client, location, scope,
        date, contract_no, revision, project_type. Only non-empty values
        overwrite the corresponding template cell — leaving a field blank
        preserves whatever the template already has. project_type drives
        the row-21 section header text (see activities_label) and, unlike
        the other header fields, has no dedicated template cell to fall
        back to — a blank/missing value defaults to "Offshore Activities".
        location and scope are also used to build the output filename
        (see ctr_tools.naming). required_info (a list of Additional
        Information items the CTR Request marked "Required") is written
        into the dedicated Comments box at cell_comments_extra (A65:G68)
        — see _write_comments_extra.

    third_party_rows: transport / hired-service lines for the "Third Party
        Activities" section, each a dict with keys comment, quantity,
        description, markup, duration, uom, rate_azn. The template has no
        such section, so one is inserted only when this list is non-empty —
        a CTR without third-party lines comes out byte-for-byte as before.

    Returns the path to the saved xlsx file.
    """
    template_path = Path(template_path)
    output_dir    = Path(output_dir)
    third_party_rows = list(third_party_rows or [])

    if not template_path.is_file():
        raise ValueError(f"AZN Template not found: {template_path}")

    try:
        output_dir.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        raise ValueError(f"Cannot create output folder {output_dir}: {e}") from e

    out_path = output_dir / ctr_output_filename(
        job_ref, "AZN", (header or {}).get("location", ""), (header or {}).get("scope", ""))
    try:
        shutil.copy2(template_path, out_path)
    except PermissionError as e:
        raise ValueError(
            f"Cannot write {out_path} — it may be open in Excel in another "
            f"window. Close it there and try again."
        ) from e
    except OSError as e:
        raise ValueError(f"Could not copy AZN template to {out_path}: {e}") from e

    sheet_name = _azn["sheet_name"]
    try:
        wb = openpyxl.load_workbook(out_path, keep_vba=False)
        ws = wb[sheet_name]
    except KeyError as e:
        raise ValueError(
            f'AZN Template {template_path} has no sheet named "{sheet_name}". '
            f"Is this the right template?"
        ) from e
    except Exception as e:
        raise ValueError(f"Could not read AZN Template {template_path}: {e}") from e

    # ── Header cells ──────────────────────────────────────────────────────────
    _write_addr(ws, _azn["cell_ctr_ref"], f"CTR-26-{job_ref} AZN")
    _write_addr(ws, _azn["cell_job_ref"], int(job_ref))

    if header:
        if header.get("client"):
            _write_addr(ws, _azn["cell_client"], header["client"])
        if header.get("sub_client"):
            _write_addr(ws, _azn["cell_sub_client"], header["sub_client"])
        if header.get("location"):
            _write_addr(ws, _azn["cell_location"], header["location"])
        if header.get("date"):
            _write_addr(ws, _azn["cell_date"], _parse_date(header["date"]))
        # Contract No and Comments are always written from the UI field,
        # even blank — unlike the other header fields here, a stale value
        # left over from the template file is actively wrong for this job
        # rather than a harmless default, so a blank UI field must blank
        # the cell instead of silently preserving whatever the template had.
        _write_addr(ws, _azn["cell_contract_no"], header.get("contract_no", ""))
        _write_comments(ws, _azn["cell_comments"], header.get("comments", ""))
        if header.get("revision") not in (None, ""):
            _write_addr(ws, _azn["cell_revision"], header["revision"])
        if header.get("scope"):
            _write_addr(ws, _azn["cell_scope"], header["scope"])

    # ── Insert extra rows before writing so overflow doesn't clobber the
    #    second section or the summary block ──────────────────────────────────
    _COLS = 9   # columns A–I used by data rows

    support_capacity = _SUPPORT_DATA_END - _SUPPORT_DATA_START + 1
    support_extra    = max(0, len(support_rows) - support_capacity)
    if support_extra:
        _insert_rows_preserving_merges(ws, _SUPPORT_DATA_END + 1, support_extra)
        for i in range(support_extra):
            _copy_row_format(ws, _SUPPORT_DATA_END, _SUPPORT_DATA_END + 1 + i, _COLS)

    # Second section has shifted down by support_extra
    other_start = _OTHER_DATA_START + support_extra
    other_end   = _OTHER_DATA_END   + support_extra

    other_capacity = other_end - other_start + 1
    other_extra    = max(0, len(other_rows) - other_capacity)
    if other_extra:
        _insert_rows_preserving_merges(ws, other_end + 1, other_extra)
        for i in range(other_extra):
            _copy_row_format(ws, other_end, other_end + 1 + i, _COLS)
    other_eff_end = other_end + other_extra

    # Total row-shift to apply to summary cell addresses
    _row_shift = support_extra + other_extra

    # ── Section-2 header text ("Onshore Activities" / "Offshore Activities"),
    #    driven by the CTR request's project type rather than fixed template
    #    text — the template ships with "Offshore Activities" hardcoded at
    #    A21/H21, which is wrong for an Onshore-only project. Written at the
    #    section's shifted title row (2 rows above its data start), to both
    #    the merged A-column cell and its standalone H-column mirror. ────────
    other_title_row = other_start - 2
    label = activities_label((header or {}).get("project_type", "") if header else "")
    ws.cell(row=other_title_row, column=1).value = label
    ws.cell(row=other_title_row, column=8).value = label

    # ── Comments box (A65:G68 by default) — the CTR Request's required
    #    Additional Information items, relocated here from the section
    #    header above (see activities_label / _write_comments_extra). ──────
    _write_comments_extra(
        ws, _azn["cell_comments_extra"], _row_shift,
        (header or {}).get("required_info") if header else None,
    )

    # ── Write section helper (captures ws via closure) ─────────────────────────
    def _write_section(rows, start, eff_end, total_gap) -> int:
        """Writes one section's rows and its total; returns the total row number."""
        clear_end = max(eff_end, start + len(rows) - 1) if rows else eff_end
        for row in range(start, clear_end + 1):
            for col in range(1, _COLS + 1):
                _set_cell(ws, row, col, None)

        for i, lr in enumerate(rows):
            r = start + i
            try:
                num_emp = float(lr.get("num_employees", 1) or 1)
            except (ValueError, TypeError):
                num_emp = 1.0
            try:
                qty = float(lr.get("quantity", 0) or 0)
            except (ValueError, TypeError):
                qty = 0.0
            try:
                rate = float(lr.get("rate_azn", 0) or 0)
            except (ValueError, TypeError):
                rate = 0.0

            ws.cell(row=r, column=1).value = strip_support_keyword(str(lr.get("comment", "")))
            ws.cell(row=r, column=2).value = num_emp
            ws.cell(row=r, column=3).value = strip_support_keyword(str(lr.get("description", "")))
            ws.cell(row=r, column=4).value = qty
            ws.cell(row=r, column=5).value = str(lr.get("uom", "Hours"))
            ws.cell(row=r, column=6).value = rate
            ws.cell(row=r, column=7).value = f"=B{r}*D{r}*F{r}"
            ws.cell(row=r, column=8).value = f"=G{r}"
            ws.cell(row=r, column=9).value = str(lr.get("nationality", "NAT"))

        # Rows left blank because there were fewer items than the section's
        # fixed capacity are hidden rather than deleted — deleting would
        # desync the SUM formula's row references below, while hiding is
        # purely cosmetic and both Excel and LibreOffice's PDF export skip
        # hidden rows when printing, so the exported PDF shows only the
        # rows that actually have data.
        #
        # Only the unused rows are contracted, and a row that does have data
        # is explicitly expanded — a template is made by saving a previous
        # CTR, so it arrives with that job's unused rows already hidden, and
        # without this a row written into one of them would be filled in
        # correctly but stay invisible in both Excel and the PDF. A section
        # that uses its whole capacity has nothing to contract and neither
        # loop does anything.
        for row in range(start, start + len(rows)):
            ws.row_dimensions[row].hidden = False
        for row in range(start + len(rows), eff_end + 1):
            ws.row_dimensions[row].hidden = True

        # The total always sits total_gap rows after the section's full
        # capacity block (eff_end), not after however many rows were
        # actually written — otherwise a request with fewer rows than the
        # template's capacity writes the total into a blank filler row
        # instead of the fixed "Total Project Support"/"Total Offshore" row.
        # Summing the whole capacity range (not just the written rows) is
        # safe — the cleared rows above are blank and contribute 0.
        total_row = eff_end + total_gap + 1
        _set_cell(ws, total_row, 7, f"=SUM(G{start}:G{eff_end})")
        # The section's own furniture — its title row, column headers and
        # total — is never spare capacity, so it is never contracted, even
        # if the template was saved with those rows hidden.
        for row in (start - 2, start - 1, *range(eff_end + 1, total_row + 1)):
            if row >= 1:
                ws.row_dimensions[row].hidden = False
        return total_row

    support_total_row = _write_section(
        support_rows, _SUPPORT_DATA_START, _SUPPORT_DATA_END + support_extra, _SUPPORT_TOTAL_GAP)
    other_total_row = _write_section(
        other_rows,   other_start,         other_eff_end,                   _OTHER_TOTAL_GAP)

    # ── Third Party Activities — inserted between the labor total and the
    #    Summary block, and only when the request actually has transport /
    #    hired-service lines (the template ships without the section). ─────
    third_party_total_row = None
    if third_party_rows:
        third_party_total_row, tp_inserted = _write_third_party_section(
            ws, other_total_row, third_party_rows, _COLS, other_start)
        _row_shift += tp_inserted

    # ── Summary cells — addresses from config, shifted by any inserted rows ────
    _r, _c = _cell_row_col(_azn["cell_summary_onshore"])
    support_summary_addr = f"{_col_letter(_c)}{_r + _row_shift}"
    ws.cell(row=_r + _row_shift, column=_c).value = f"=G{support_total_row}"

    _r, _c = _cell_row_col(_azn["cell_summary_offshore"])
    other_summary_row  = _r + _row_shift
    other_summary_addr = f"{_col_letter(_c)}{other_summary_row}"
    ws.cell(row=other_summary_row, column=_c).value = f"=G{other_total_row}"

    _r, _c = _cell_row_col(_azn["cell_summary_combined"])
    combined_row = _r + _row_shift
    summary_addrs = [support_summary_addr, other_summary_addr]

    # The Summary block needs its own line for the third-party total, added
    # right below the labor one; every numbered item below it (Contingency)
    # shifts down a place and is renumbered to match.
    summary_extra = 0
    if third_party_total_row is not None:
        tp_summary_row = other_summary_row + 1
        combined_row  += 1
        summary_extra  = 1
        _insert_rows_preserving_merges(ws, tp_summary_row, 1)
        _copy_row_format(ws, other_summary_row, tp_summary_row, _COLS)
        ws.merge_cells(start_row=tp_summary_row, start_column=2,
                       end_row=tp_summary_row, end_column=5)
        ws.cell(row=tp_summary_row, column=2).value = _azn.get(
            "third_party_summary_label", "Total Other Activities")
        ws.cell(row=tp_summary_row, column=_c).value = f"=G{third_party_total_row}"
        summary_addrs.append(f"{_col_letter(_c)}{tp_summary_row}")

        item_no = ws.cell(row=other_summary_row, column=1).value
        if isinstance(item_no, (int, float)) and not isinstance(item_no, bool):
            for row in range(tp_summary_row, combined_row):
                existing = ws.cell(row=row, column=1).value
                if row == tp_summary_row or (
                    isinstance(existing, (int, float)) and not isinstance(existing, bool)
                ):
                    item_no = int(item_no) + 1
                    ws.cell(row=row, column=1).value = item_no

    ws.cell(row=combined_row, column=_c).value = "=" + "+".join(summary_addrs)

    # Every row inserted above moved the document's tail down by that much;
    # the print area has to follow or the bottom of the CTR silently stops
    # appearing in the PDF.
    _extend_print_area(ws, _row_shift + summary_extra)
    _clear_placeholder_cells(ws, _azn.get("placeholder_texts", []))
    _fit_to_page_width(ws)

    try:
        wb.save(out_path)
    except PermissionError as e:
        raise ValueError(
            f"Cannot save {out_path} — it may be open in Excel in another "
            f"window. Close it there and try again."
        ) from e
    finally:
        wb.close()
    return out_path

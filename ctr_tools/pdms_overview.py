"""
ctr_tools/pdms_overview.py

The "PDMS overview" sheet of the CTR statistics workbook: the same figures as the PDMS proposal
(Word), laid out like the PDMS request's own pricing sheet — labour, equipment, consumables with
the transport lines, and the summary — but with the arithmetic left in the cells, so a rate or a
quantity can be changed and the totals follow. Unlike the (client-facing) Word document it also
shows each line's Stock Code, and the Encoding of a manpower line.

Every price keeps its own currency; a CBAR-rate block at the top holds the conversion rates, how
each was last checked or changed, and each line looks its rate up there (VLOOKUP on the currency
code). Amounts in manats use the sheet's own number format. The figures come from
builder_pdms.compute_pdms(), so the sheet and the Word document cannot disagree.

PROVISIONAL (task log MA202): which rate / date PDMS converts with is not confirmed yet.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ctr_tools.builder_pdms import letter_date, pdms_defaults

SHEET_NAME = "PDMS overview"
_AZN = '#,##0.00\\ [$₼-42C]'          # the manat format of the PDMS request's own sheet
_NUM = '#,##0.00'
_FX = '0.0000'
_THIN = Side(style="thin", color="999999")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HDR = {"labour": "D9D9D9", "equipment": "BFBFBF", "consumables": "D0CECE", "transport": "D0CECE"}
_ROW, _ROW2, _TOTAL = "DCE6F1", "D9E2F3", "8EAADB"
_WIDTHS = (2, 40, 17, 13, 13, 13, 15, 15, 15, 17, 17, 19, 15, 15)      # A … N


def _fill(rgb: str) -> PatternFill:
    return PatternFill("solid", start_color=rgb, end_color=rgb)


class _Sheet:
    """Styled cells at (row, column) of a worksheet."""

    def __init__(self, ws):
        self.ws = ws

    def put(self, row, col, value, *, bold=False, fill=None, fmt=None, align=None, border=True,
            color=None, italic=False, size=None):
        c = self.ws.cell(row=row, column=col, value=value)
        c.font = Font(name="Arial", size=size or 10, bold=bold, italic=italic, color=color)
        if fill:
            c.fill = _fill(fill)
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(horizontal=align, vertical="center") if align else Alignment(vertical="center")
        if border:
            c.border = _BOX
        return c

    def title(self, row, text):
        self.put(row, 2, text, bold=True, size=11, border=False)


class _Table:
    """A table whose columns are named, so a formula can say {qty} instead of a letter that moves
    when a column is added. Columns start at B; `at(key, row)` is the cell address."""

    def __init__(self, sh: _Sheet, kind: str, columns: list[tuple[str, str]]):
        self.sh, self.kind = sh, kind
        self.keys = [k for k, _h in columns]
        self.headers = [h for _k, h in columns]

    def col(self, key: str) -> int:
        return 2 + self.keys.index(key)

    def letter(self, key: str) -> str:
        return get_column_letter(self.col(key))

    def at(self, key: str, row: int) -> str:
        return f"{self.letter(key)}{row}"

    def f(self, template: str, row: int) -> str:
        """"=ROUND({qty}*{unit},2)" with every {key} turned into that column's cell in `row`."""
        return template.format(**{k: self.at(k, row) for k in self.keys})

    def header(self, row: int) -> None:
        for i, text in enumerate(self.headers):
            self.sh.put(row, 2 + i, text, bold=True, fill=_HDR[self.kind], align="center")

    def put(self, row: int, key: str, value, **kw):
        return self.sh.put(row, self.col(key), value, **kw)

    def total_row(self, row: int, label: str, label_key: str, total_key: str, formula):
        """A dark total row: everything filled, the label right-aligned in `label_key`'s column, the
        figure in `total_key`'s."""
        for c in range(2, self.col(total_key) + 1):
            self.sh.put(row, c, None, fill=_TOTAL)
        self.sh.put(row, self.col(label_key), label, bold=True, fill=_TOTAL, align="right")
        return self.put(row, total_key, formula, bold=True, fill=_TOTAL, fmt=_AZN)


def add_overview_sheet(wb, data: dict, header: dict, rate_days: dict | None = None,
                       index: int = 0) -> dict:
    """
    Add the "PDMS overview" sheet to workbook `wb` (at position `index`) and return the addresses
    of its totals (for tests / callers): {"labour_total", "equipment_total", "materials_total",
    "grand_total"} as cell addresses (equipment / materials only when there is such a section),
    plus "rates" (the CBAR block's range).

    data: builder_pdms.compute_pdms() result.
    header: the proposal's header fields (prp_number, date, customer_name, attention_name, scope,
        project_type …).
    rate_days: {currency code: {"day": CBAR day "YYYY-MM-DD" or None when typed, "stamp": when the rate
        was last checked or changed, "action": what did it} for the Source and history columns (a plain
        {code: day} also works).
    """
    cfg = pdms_defaults()
    ws = wb.create_sheet(SHEET_NAME, index)
    ws.sheet_view.showGridLines = False
    sh = _Sheet(ws)
    for i, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    sh.put(1, 2, f"PDMS overview — {header.get('prp_number', '')}", bold=True, size=14, border=False)
    sh.put(2, 2, "The same figures as the PDMS proposal (Word), with the arithmetic in the cells: change a "
                 "rate or a quantity and the totals follow. Prices are in their own currency and are "
                 "converted to manats with the CBAR rates below (provisional rule); a total is rounded to "
                 "2 decimals. Stock codes and Encodings are shown here only, not in the proposal.",
           italic=True, color="666666", border=False)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=12)
    ws.cell(row=2, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30

    # ── the job ──
    info = [
        ("PRP No", header.get("prp_number", "")),
        ("Date", letter_date(header.get("date", ""))),
        ("Customer", header.get("customer_name") or cfg["customer_name"]),
        ("Attention", header.get("attention_name", "")),
        ("Job", header.get("scope", "")),
        ("Location", data["location"]),
    ]
    for i, (k, v) in enumerate(info):
        sh.put(4 + i, 2, k, bold=True, fill="F2F2F2")
        sh.put(4 + i, 3, v)
        ws.merge_cells(start_row=4 + i, start_column=3, end_row=4 + i, end_column=6)

    # ── CBAR rates (AZN is always 1): the rate, where it is from, when it was last checked or changed, by what ──
    sh.put(3, 8, "CBAR rates", bold=True, size=11, border=False)
    for c, text in zip((8, 9, 10, 11, 12), ("Currency", "AZN per 1", "Source", "Last checked / changed", "Last action")):
        sh.put(4, c, text, bold=True, fill=_HDR["labour"], align="center")
    for c in (13, 14):
        sh.put(4, c, None, fill=_HDR["labour"])
    ws.merge_cells(start_row=4, start_column=12, end_row=4, end_column=14)
    ordered = ["AZN"] + [c for c in sorted(data["rates"]) if c != "AZN"]
    for i, code in enumerate(ordered):
        r = 5 + i
        sh.put(r, 8, code, align="center")
        if code == "AZN":
            sh.put(r, 9, 1, fmt=_FX)
            sh.put(r, 10, "fixed")
            sh.put(r, 11, None)
            sh.put(r, 12, "—")
        else:
            rate = data["rates"].get(code)
            sh.put(r, 9, rate if rate else None, fmt=_FX)
            ev = (rate_days or {}).get(code)
            ev = ev if isinstance(ev, dict) else {"day": ev, "stamp": "", "action": ""}
            day = ev.get("day")
            sh.put(r, 10, f"CBAR {day}" if day else ("typed by hand" if rate else "—"))
            sh.put(r, 11, ev.get("stamp") or "—", align="center")
            sh.put(r, 12, ev.get("action") or "—")
        for c in (13, 14):
            sh.put(r, c, None)
        ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=14)
    rate_first, rate_last = 5, 4 + len(ordered)
    rates_rng = f"$H${rate_first}:$I${rate_last}"

    def fx(table: _Table, row: int, cur_key: str) -> str:
        return f'=IFERROR(VLOOKUP({table.at(cur_key, row)},{rates_rng},2,FALSE),0)'

    row = max(11, rate_last + 2)
    out: dict = {"rates": rates_rng}

    # ── Table 1: labour (with the Encoding of each line) ──
    t = _Table(sh, "labour", [
        ("desc", "DESCRIPTION"), ("enc", "Encoding"), ("nat", "Nationality"), ("loc", "Location"), ("days", "Days"),
        ("men", "No of men"), ("native", "Normal day rate (own currency)"), ("cur", "Cur."), ("fx", "CBAR rate"),
        ("rate", "Normal Day Rate (AZN)"), ("total", "Total (AZN)")])
    sh.title(row, "Table 1 — Labour rates")
    row += 1
    t.header(row)
    row += 1
    first = row
    for i, ln in enumerate(data["labour"]):
        fill = _ROW if i % 2 == 0 else _ROW2
        t.put(row, "desc", ln["description"], fill=fill)
        t.put(row, "enc", ln.get("encoding") or None, fill=fill, align="center")
        t.put(row, "nat", ln["nationality"], fill=fill, align="center")
        t.put(row, "loc", data["location"], fill=fill, align="center")
        t.put(row, "days", ln["days"], fill=fill, align="center", fmt="General")
        t.put(row, "men", ln["men"], fill=fill, align="center", fmt="General")
        t.put(row, "native", ln["native"], fill=fill, fmt=_NUM)
        t.put(row, "cur", ln["cur"], fill=fill, align="center")
        t.put(row, "fx", fx(t, row, "cur"), fill=fill, fmt=_FX)
        t.put(row, "rate", t.f("={native}*{fx}", row), fill=_ROW2, fmt=_AZN)
        t.put(row, "total", t.f("=ROUND({days}*{men}*{rate},2)", row), fill=_ROW2, fmt=_AZN)
        row += 1
    hb = data["hotbed"]
    if hb is not None:            # always there once it is ticked; calculated only when it is complete
        for key in ("enc", "nat", "loc", "men"):
            t.put(row, key, None, fill=_ROW)
        t.put(row, "desc", cfg["hotbed_label"], fill=_ROW)
        t.put(row, "days", hb.get("days"), fill=_ROW, align="center", fmt="General")
        t.put(row, "native", hb.get("native"), fill=_ROW, fmt=_NUM)
        t.put(row, "cur", hb.get("cur", "AZN") if hb.get("rate") else None, fill=_ROW, align="center")
        t.put(row, "fx", fx(t, row, "cur") if hb.get("rate") else None, fill=_ROW, fmt=_FX)
        t.put(row, "rate", t.f("={native}*{fx}", row) if hb.get("rate") else None, fill=_ROW, fmt=_AZN)
        t.put(row, "total", t.f('=IF(AND(ISNUMBER({days}),ISNUMBER({rate})),ROUND({days}*{rate},2),"")', row),
              fill=_ROW, fmt=_AZN)
        row += 1
    last = row - 1
    tc = t.letter("total")
    t.total_row(row, "LABOUR TOTAL:", "desc", "total", f"=SUM({tc}{first}:{tc}{last})" if last >= first else 0)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=t.col("rate"))
    out["labour_total"] = t.at("total", row)
    labour_total_cell = out["labour_total"]
    row += 2

    # ── Table 2: equipment (with its Stock Code) ──
    equip_total_cell = None
    if data["equipment"]:
        t = _Table(sh, "equipment", [
            ("desc", "Description"), ("code", "Stock code"), ("qty", "QTY"), ("days", "Duration (days)"),
            ("uom", "UOM"), ("native", "Unit price (own currency)"), ("cur", "Cur."), ("fx", "CBAR rate"),
            ("unit", "Unit price (AZN)"), ("total", "Total (AZN)")])
        sh.title(row, "Table 2 — Equipment rates")
        row += 1
        t.header(row)
        row += 1
        first = row
        for ln in data["equipment"]:
            t.put(row, "desc", ln["description"], fill=_ROW2)
            t.put(row, "code", ln.get("code") or None, fill=_ROW2, align="center")
            t.put(row, "qty", ln["qty"], fill=_ROW2, align="center", fmt="General")
            t.put(row, "days", ln["days"], fill=_ROW2, align="center", fmt="General")
            t.put(row, "uom", ln["uom"], fill=_ROW2, align="center")
            t.put(row, "native", ln["native"], fill=_ROW2, fmt=_NUM)
            t.put(row, "cur", ln["cur"], fill=_ROW2, align="center")
            t.put(row, "fx", fx(t, row, "cur"), fill=_ROW2, fmt=_FX)
            t.put(row, "unit", t.f("={native}*{fx}", row), fill=_ROW2, fmt=_AZN)
            t.put(row, "total", t.f("=ROUND({qty}*{days}*{unit},2)", row), fill=_ROW2, fmt=_AZN)
            row += 1
        tc = t.letter("total")
        t.total_row(row, "TOTAL:", "unit", "total", f"=SUM({tc}{first}:{tc}{row - 1})")
        equip_total_cell = out["equipment_total"] = t.at("total", row)
        row += 2

    # ── Table 3: consumables (with their Stock Code), then the transport lines ──
    cons_cells: list[str] = []
    n_table = 2 + (1 if data["equipment"] else 0)
    if data["consumables"]:
        t = _Table(sh, "consumables", [
            ("desc", "Material Description"), ("code", "Stock code"), ("uom", "UOM"), ("qty", "Q-ty"),
            ("native", "Cost (own currency)"), ("cur", "Cur."), ("fx", "CBAR rate"), ("markup", "Mark-up"),
            ("unit", "Unit cost incl. mark-up (AZN)"), ("total", "Total (AZN)")])
        sh.title(row, f"Table {n_table} — Consumables costs")
        row += 1
        t.header(row)
        row += 1
        first = row
        for ln in data["consumables"]:
            t.put(row, "desc", ln["description"], fill=_ROW2)
            t.put(row, "code", ln.get("code") or None, fill=_ROW2, align="center")
            t.put(row, "uom", ln["uom"], fill=_ROW2, align="center")
            t.put(row, "qty", ln["qty"], fill=_ROW2, align="center", fmt="General")
            t.put(row, "native", ln["native"], fill=_ROW2, fmt=_NUM)
            t.put(row, "cur", ln["cur"], fill=_ROW2, align="center")
            t.put(row, "fx", fx(t, row, "cur"), fill=_ROW2, fmt=_FX)
            t.put(row, "markup", ln["markup"], fill=_ROW2, fmt="0.0%")
            t.put(row, "unit", t.f("={native}*{fx}*(1+{markup})", row), fill=_ROW2, fmt=_AZN)
            t.put(row, "total", t.f("=ROUND({qty}*{unit},2)", row), fill=_ROW2, fmt=_AZN)
            row += 1
        tc = t.letter("total")
        cons_cells.append(f"SUM({tc}{first}:{tc}{row - 1})")
        row += 1
    if data["transport"]:
        t = _Table(sh, "transport", [
            ("desc", "Description"), ("code", ""), ("qty", "Qty"), ("dur", "Duration"), ("uom", "UOM"),
            ("native", "Rate (own currency)"), ("cur", "Cur."), ("fx", "CBAR rate"), ("markup", "Mark-up"),
            ("total", "Total (AZN)")])
        sh.title(row, "Transport (listed at the end of the consumables table in the proposal)")
        row += 1
        t.header(row)
        row += 1
        first = row
        for ln in data["transport"]:
            t.put(row, "desc", ln["description"], fill=_ROW2)
            t.put(row, "code", None, fill=_ROW2)                       # transport has no stock code
            t.put(row, "qty", ln["qty"], fill=_ROW2, align="center", fmt="General")
            t.put(row, "dur", ln["duration"], fill=_ROW2, align="center", fmt="General")
            t.put(row, "uom", ln["uom"], fill=_ROW2, align="center")
            t.put(row, "native", ln["native"], fill=_ROW2, fmt=_NUM)
            t.put(row, "cur", ln["cur"], fill=_ROW2, align="center")
            t.put(row, "fx", fx(t, row, "cur"), fill=_ROW2, fmt=_FX)
            t.put(row, "markup", ln["markup"], fill=_ROW2, fmt="0.0%")
            t.put(row, "total", t.f("=ROUND({qty}*{native}*{fx}*{dur}*(1+{markup}),2)", row), fill=_ROW2, fmt=_AZN)
            row += 1
        tc = t.letter("total")
        cons_cells.append(f"SUM({tc}{first}:{tc}{row - 1})")
        row += 1
    materials_cell = None
    if cons_cells:
        tot_col = 11                                                    # K: both tables end here
        for c in range(2, tot_col + 1):
            sh.put(row, c, None, fill=_TOTAL)
        sh.put(row, 2, "MATERIALS TOTAL:", bold=True, fill=_TOTAL, align="right")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=tot_col - 1)
        sh.put(row, tot_col, "=" + "+".join(cons_cells), bold=True, fill=_TOTAL, fmt=_AZN)
        materials_cell = out["materials_total"] = f"{get_column_letter(tot_col)}{row}"
        row += 2

    # ── Summary ──
    sh.title(row, "Summary")
    row += 1
    summary = [("LABOUR:", f"={labour_total_cell}")]
    if equip_total_cell:
        summary.append(("MATERIAL:", f"={equip_total_cell}"))
    if materials_cell:
        summary.append(("CONSUMABLES:", f"={materials_cell}"))
    first = row
    for label, formula in summary:
        sh.put(row, 2, label, align="right")
        sh.put(row, 3, formula, bold=True, fmt=_AZN)
        row += 1
    sh.put(row, 2, "GRAND TOTAL:", bold=True, fill=_TOTAL, align="right")
    sh.put(row, 3, f"=SUM(C{first}:C{row - 1})", bold=True, fill=_TOTAL, fmt=_AZN)
    out["grand_total"] = f"C{row}"
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = "1976D2"
    return out

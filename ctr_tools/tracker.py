"""
ctr_tools/tracker.py

Backend for the CTR Tracker tab: reads header fields off a generated CTR
output file (AZN or USD template — both share the same header cell layout,
see "ctr_extract" in template_config.json), the site/location lookup, and
the pre-write backup helper. The actual write into the tracker workbook
lives in tracker_fast.py.

Column positions for reading come from CFG so they can be adjusted without
touching this file when the layout changes.
"""

from __future__ import annotations

import logging
import shutil
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as _col_idx

from ctr_tools.config import CFG

log = logging.getLogger(__name__)


@dataclass
class CTREntry:
    """One row queued for the tracker. All fields are user-editable in the
    UI after extraction, so nothing here is assumed authoritative."""
    source_path:  str = ""
    client:       str = ""
    ctr_number:   str = ""
    ctr_date:     date | None = None
    revision:     str = ""          # blank ("") means "none"
    description:  str = ""
    value:        float | None = None
    currency:     str = ""
    site_hint:    str = ""          # raw site text read from the CTR file (cell B4)
    site:         str = ""          # key into ctr_tracker_locations, once resolved/confirmed
    project_code: str = ""
    tracker_location: str = ""      # "Onshore" / "Offshore" / "Georgia"
    job_type:     str = ""          # one of ctr_tracker_job_types — user-picked, not read from the CTR file
    value_is_estimate: bool = False  # True if the value cell had no cached
                                      # number and had to be left for manual entry

    # Cost-breakdown fields — best-effort, "not all applicable" to every
    # CTR (a field with no matching line in the file stays None and is
    # simply not written). See CFG["ctr_tracker"]'s col_labor/col_equipment/
    # etc. for where each lands in the tracker.
    activity_type:                str = ""          # "Onshore" or "Offshore", from the CTR's own labor section label
    total_project_support:        float | None = None
    total_activities:              float | None = None  # the activity_type-matching "Total X Activities" figure
    total_third_party:            float | None = None
    total_equipment:               float | None = None
    general_consumables:          float | None = None
    consumables_markup_amount:    float | None = None
    transportation_customs:       float | None = None
    transportation_markup_amount: float | None = None


def location_options() -> dict[str, dict]:
    """Site name -> {project_code, tracker_location}, from config."""
    return CFG.get("ctr_tracker_locations", {})


def job_type_options() -> list[str]:
    """Choices offered in the Job Type dropdown, from config."""
    return CFG.get("ctr_tracker_job_types", [])


def match_site(site_hint: str) -> str | None:
    """Case-insensitive match of a CTR file's raw site text (cell B4)
    against the configured site list. Returns the matching site name (the
    dict key to use with location_options()), or None if nothing matches —
    the caller should fall back to asking the user to pick one."""
    hint = site_hint.strip().lower()
    if not hint:
        return None
    for name in location_options():
        if name.strip().lower() == hint:
            return name
    return None


def company_project_code(client: str) -> str | None:
    """Case-insensitive match of a CTR's Client field against
    ctr_tracker_companies. Some companies (e.g. Turan Drilling, CDC) always
    use the same Project Code no matter which site/location the CTR is
    for — this takes priority over the location-based lookup above when it
    matches. Returns None if the client isn't one of the configured
    companies, so the caller falls back to location_options()."""
    client_norm = client.strip().lower()
    if not client_norm:
        return None
    for name, code in CFG.get("ctr_tracker_companies", {}).items():
        if name.strip().lower() == client_norm:
            return code
    return None


def _main_sheet(wb) -> "openpyxl.worksheet.worksheet.Worksheet":
    azn_name = CFG["azn_template"]["sheet_name"]
    usd_name = CFG["usd_template"]["main_sheet"]
    if azn_name in wb.sheetnames:
        return wb[azn_name]
    if usd_name in wb.sheetnames:
        return wb[usd_name]
    log.warning(
        "Neither the AZN main sheet (%r) nor the USD main sheet (%r) was "
        "found among %r — falling back to the first sheet (%r), which may "
        "not carry the expected header layout and could read blank/wrong "
        "values.", azn_name, usd_name, wb.sheetnames, wb.sheetnames[0],
    )
    return wb[wb.sheetnames[0]]


def _find_value(ws_calc, ws_raw, label: str, label_cols: list[str],
                 result_col: str, search_max_row: int,
                 exact: bool = False) -> tuple[float | None, bool]:
    """
    Search label_cols for a cell matching label (case-insensitive; prefix
    match by default, exact match when exact=True — labels like
    "Estimated CTR Total " carry template-inconsistent trailing text) and
    read the total from result_col — the summary section's row shifts
    depending on how many line items a CTR has, so a fixed cell address
    isn't reliable. A label that appears more than once (its own section's
    subtotal, then again in the final Summary rollup with the same figure)
    resolves to the last match, which lands on the Summary rollup when
    there is one.

    Returns (value, is_estimate). is_estimate is True when the label was
    found but the result cell has no cached number (the formula was never
    recalculated by Excel), so the caller should prompt for manual entry.
    """
    label_norm = label.strip().lower()
    label_col_idxs = [_col_idx(c) for c in label_cols]
    result_col_idx = _col_idx(result_col)

    found_row = None
    for row in range(1, search_max_row + 1):
        for col in label_col_idxs:
            cell_val = ws_calc.cell(row=row, column=col).value
            if not isinstance(cell_val, str):
                continue
            text = cell_val.strip().lower()
            if text == label_norm or (not exact and text.startswith(label_norm)):
                found_row = row
                break

    if found_row is None:
        return None, False

    result = ws_calc.cell(row=found_row, column=result_col_idx).value
    if isinstance(result, (int, float)):
        return float(result), False
    # No cached value — fall back to a literal (non-formula) number in the
    # un-recalculated workbook, otherwise give up gracefully.
    raw = ws_raw.cell(row=found_row, column=result_col_idx).value
    if isinstance(raw, (int, float)):
        return float(raw), False
    return None, True


def _find_activity_type(ws_calc, cfg: dict) -> str:
    """The manpower section's own "Onshore Activities" / "Offshore
    Activities" label, minus the " Activities" suffix. Empty string if not
    found (e.g. a CTR type without a manpower section).

    Matched on what the cell *starts with*, not on the whole cell: CTRs
    generated before builder_azn.activities_label moved required Additional
    Information into its own Comments box instead had it appended right
    onto this header ("Offshore Activities : Per Diem required"), and an
    exact-match test would read one of those older files as no manpower
    section at all."""
    labels = {label.strip().lower(): label for label in cfg["activity_labels"]}
    col = _col_idx(cfg["activity_label_col"])
    for row in range(1, cfg["search_max_row"] + 1):
        cell_val = ws_calc.cell(row=row, column=col).value
        if not isinstance(cell_val, str):
            continue
        text = cell_val.strip().lower()
        for key, matched in labels.items():
            if text.startswith(key):
                return matched.rsplit(" ", 1)[0]   # "Offshore Activities" -> "Offshore"
    return ""


def extract_ctr_data(path: str | Path) -> CTREntry:
    """Read client/CTR number/date/revision/description/value/currency off
    a generated CTR output file. Fields that can't be read are left blank
    for the user to fill in by hand."""
    path = Path(path)
    cfg = CFG["ctr_extract"]

    wb_calc = openpyxl.load_workbook(path, data_only=True)
    ws_calc = _main_sheet(wb_calc)
    wb_raw  = openpyxl.load_workbook(path, data_only=False)
    ws_raw  = wb_raw[ws_calc.title]

    client = str(ws_calc[cfg["cell_client"]].value or "").strip()

    ctr_number = str(ws_calc[cfg["cell_ctr_ref"]].value or "").strip()
    site_hint = str(ws_calc[cfg["cell_location"]].value or "").strip()

    raw_date = ws_calc[cfg["cell_date"]].value
    ctr_date = raw_date.date() if isinstance(raw_date, datetime) else (
        raw_date if isinstance(raw_date, date) else None
    )

    raw_rev = ws_calc[cfg["cell_revision"]].value
    revision = "" if raw_rev in (None, "", 0) else str(raw_rev).strip()

    description = str(ws_calc[cfg["cell_scope"]].value or "").strip()
    currency = str(ws_calc[cfg["cell_currency"]].value or "").strip().upper()

    def _val(label: str, exact: bool = False) -> float | None:
        result, _ = _find_value(
            ws_calc, ws_raw, label, cfg["label_cols"], cfg["value_result_col"],
            cfg["search_max_row"], exact=exact,
        )
        return result

    value, is_estimate = _find_value(
        ws_calc, ws_raw, cfg["value_label"], [cfg["value_label_col"]],
        cfg["value_result_col"], cfg["search_max_row"],
    )

    activity_type = _find_activity_type(ws_calc, cfg)
    total_activities = _val(f"Total {activity_type} Activities") if activity_type else None

    wb_calc.close()
    wb_raw.close()

    return CTREntry(
        source_path=str(path),
        client=client,
        ctr_number=ctr_number,
        site_hint=site_hint,
        ctr_date=ctr_date,
        revision=revision,
        description=description,
        value=value,
        currency=currency,
        value_is_estimate=is_estimate,
        activity_type=activity_type,
        total_project_support=_val(cfg["label_total_project_support"]),
        total_activities=total_activities,
        total_third_party=_val(cfg["label_total_third_party"]),
        total_equipment=_val(cfg["label_total_equipment"]),
        general_consumables=_val(cfg["label_general_consumables"]),
        consumables_markup_amount=_val(cfg["label_consumables_markup"]),
        transportation_customs=_val(cfg["label_transportation_customs"]),
        transportation_markup_amount=_val(cfg["label_transportation_markup"]),
    )


def backup_tracker(tracker_path: str | Path) -> Path:
    """Copies the tracker workbook to a timestamped backup next to it before
    it gets overwritten — this is a shared, macro-driven master file, so a
    saved-over mistake shouldn't be unrecoverable. Returns the backup path."""
    tracker_path = Path(tracker_path)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_path = tracker_path.with_name(
        f"{tracker_path.stem}.backup-{stamp}{tracker_path.suffix}"
    )
    shutil.copy2(tracker_path, backup_path)
    return backup_path

"""
ctr_generator/builder_azn.py

Copies the AZN CTR template and writes labor rows into sheet "Main AZN".

Row/column positions are read from ctr_generator/template_config.json
(section "azn_template") so they can be adjusted without touching this
file when a new template version shifts the layout.

Default layout (217_AZN template):
  Row 3  : E3 = CTR ref string, O3 = job ref integer
  Row 7  : "Project Support" section header (fixed text)
  Row 8  : column headers for the Project Support section
  Rows 9–18 : "support" section data rows — any manpower row whose
    description contains the keyword "support" (case-insensitive; see
    is_support_manpower), regardless of its Onshore/Offshore type
  Row 20 : G20 = total project support
  Row 21 : section header for every other manpower row — text is
    "Onshore Activities" or "Offshore Activities" depending on the CTR
    request's project type (see activities_label), not fixed template text
  Row 22 : column headers for that section
  Rows 23–62: data rows for every manpower row NOT matched as "support"
  Row 64 : G64 = total for that section
  Summary (rows 66–71):
    G67 = total project support
    G68 = total of the other section
    G71 = estimated CTR total

  Editable header fields (rows 3–6, shared layout with USD template):
    B3 = Client, C3 = Sub-Client, B4 = Location, E4 = Date,
    G4 = Contract No, E5 = Revision, A6 = Scope / description (merged A6:G6)

Input values (comment, employees, quantity, rate, etc.) are written as
plain Python values. Anything derived from another cell — row totals,
section totals, summary cells — is written as an Excel formula referencing
that cell, so editing an input in Excel recalculates its dependents. This
only matters for the .xlsx output; the PDF export is a flat rendering, so
formulas there simply show their last-calculated value.

The "support" keyword itself (see is_support_manpower / strip_support_keyword)
is an internal routing marker, not client-facing text — it's stripped from
the Comment and Description cells of every written row before they reach
the document, so "Painter Support" appears as plain "Painter".
"""

from __future__ import annotations

import re
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as _col_idx
from openpyxl.utils import get_column_letter as _col_letter

from ctr_generator.config import CFG
from ctr_generator.naming import ctr_output_filename

_azn = CFG["azn_template"]

_SUPPORT_DATA_START = _azn["onshore_data_start"]
_SUPPORT_DATA_END   = _azn["onshore_data_end"]
_SUPPORT_TOTAL_GAP  = _azn["onshore_total_gap"]   # blank rows between data end and total

_OTHER_DATA_START = _azn["offshore_data_start"]
_OTHER_DATA_END   = _azn["offshore_data_end"]
_OTHER_TOTAL_GAP  = _azn["offshore_total_gap"]

_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y")

# Manpower rows whose description contains this keyword (any casing —
# "Support", "SUPPORT", "Supporting", "Life Support", etc. all match as a
# plain substring) go into the fixed "Project Support" rows (9–18); every
# other manpower row goes into the section starting at row 23.
_SUPPORT_KEYWORD = "support"
_SUPPORT_RE = re.compile(re.escape(_SUPPORT_KEYWORD), re.IGNORECASE)


def is_support_manpower(description: str) -> bool:
    """True if a manpower description should be treated as "support" manpower."""
    return _SUPPORT_KEYWORD in (description or "").lower()


def strip_support_keyword(text: str) -> str:
    """
    Removes the "support" keyword from a manpower match string before it's
    looked up in the pricebook — a role like "Scaffolder Support" should
    match the base pricebook item "Scaffolder" rather than fail as an
    unmatched exact string. Cleans up whitespace/punctuation left behind by
    the removal (e.g. "Painter (Support)" → "Painter"); falls back to the
    original text if stripping would leave nothing.
    """
    if not text:
        return text
    stripped = _SUPPORT_RE.sub("", text)
    stripped = re.sub(r"\(\s*\)", "", stripped)
    stripped = re.sub(r"\s{2,}", " ", stripped).strip(" -_/,()")
    return stripped or text.strip()


def activities_label(project_type: str) -> str:
    """
    Section header text for the non-"support" manpower section (row 21/H21):
    "Onshore Activities" or "Offshore Activities", chosen from the CTR
    request's project type. Falls back to "Offshore Activities" (the
    template's original hardcoded text) when project_type is blank or
    doesn't recognize either word.
    """
    if "onshore" in (project_type or "").lower():
        return _azn.get("label_onshore_activities", "Onshore Activities")
    return _azn.get("label_offshore_activities", "Offshore Activities")


def _cell_row_col(addr: str) -> tuple[int, int]:
    """Convert a cell address like 'G67' to (row=67, col=7)."""
    col_str = "".join(c for c in addr if c.isalpha())
    row_str = "".join(c for c in addr if c.isdigit())
    return int(row_str), _col_idx(col_str)


def _insert_rows_preserving_merges(ws, insert_row: int, amount: int) -> None:
    """
    openpyxl's insert_rows() shifts cell values down but leaves merged-cell
    ranges at their original row numbers — a stale merge left behind at the
    insertion point (e.g. a section-header row) silently swallows any data
    later written into whichever row now occupies that position. Unmerging
    *after* insert_rows() also fails (the merge's cached cell objects no
    longer match reality), so affected ranges must be unmerged first,
    inserted around, then re-merged at their shifted position.
    """
    affected = [
        (mr.min_row, mr.min_col, mr.max_row, mr.max_col)
        for mr in ws.merged_cells.ranges if mr.min_row >= insert_row
    ]
    for min_row, min_col, max_row, max_col in affected:
        ws.unmerge_cells(start_row=min_row, start_column=min_col,
                          end_row=max_row, end_column=max_col)

    ws.insert_rows(insert_row, amount)

    for min_row, min_col, max_row, max_col in affected:
        ws.merge_cells(start_row=min_row + amount, start_column=min_col,
                        end_row=max_row + amount, end_column=max_col)


def _copy_row_format(ws, src_row: int, dst_row: int, col_count: int) -> None:
    """Copy cell formatting (borders, fill, font, alignment) from one row to another."""
    src_dim = ws.row_dimensions.get(src_row)
    if src_dim and src_dim.height:
        ws.row_dimensions[dst_row].height = src_dim.height
    for col in range(1, col_count + 1):
        src = ws.cell(row=src_row, column=col)
        dst = ws.cell(row=dst_row, column=col)
        if src.has_style:
            dst.font         = copy(src.font)
            dst.border       = copy(src.border)
            dst.fill         = copy(src.fill)
            dst.number_format = src.number_format
            dst.alignment    = copy(src.alignment)


def _parse_date(value):
    """Best-effort string → datetime conversion; returns value unchanged if unparseable."""
    if isinstance(value, str):
        for fmt in _DATE_FORMATS:
            try:
                return datetime.strptime(value, fmt)
            except ValueError:
                continue
    return value


def build_azn(
    support_rows: list[dict],
    other_rows: list[dict],
    template_path: str | Path,
    output_dir: str | Path,
    job_ref: str,
    header: dict | None = None,
) -> Path:
    """
    Write labor rows into a copy of the AZN template.

    support_rows: manpower rows whose description matched the "support"
        keyword (see is_support_manpower) — written into rows 9–18.
    other_rows: every other manpower row — written into rows 23+.

    Each row is a dict with keys:
        comment, num_employees (int/float), description (str),
        quantity (float), uom (str), rate_azn (float), nationality (str)

    header: optional dict with keys client, sub_client, location, scope,
        date, contract_no, revision, project_type. Only non-empty values
        overwrite the corresponding template cell — leaving a field blank
        preserves whatever the template already has. project_type drives
        the row-21 section header text (see activities_label) and, unlike
        the other header fields, has no dedicated template cell to fall
        back to — a blank/missing value defaults to "Offshore Activities".
        location and scope are also used to build the output filename
        (see ctr_generator.naming).

    Returns the path to the saved xlsx file.
    """
    template_path = Path(template_path)
    output_dir    = Path(output_dir)

    if not template_path.is_file():
        raise ValueError(f"AZN Template not found: {template_path}")

    try:
        output_dir.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        raise ValueError(f"Cannot create output folder {output_dir}: {e}") from e

    out_path = output_dir / ctr_output_filename(
        job_ref, "AZN", (header or {}).get("location", ""), (header or {}).get("scope", ""))
    try:
        shutil.copy2(template_path, out_path)
    except PermissionError as e:
        raise ValueError(
            f"Cannot write {out_path} — it may be open in Excel in another "
            f"window. Close it there and try again."
        ) from e
    except OSError as e:
        raise ValueError(f"Could not copy AZN template to {out_path}: {e}") from e

    sheet_name = _azn["sheet_name"]
    try:
        wb = openpyxl.load_workbook(out_path, keep_vba=False)
        ws = wb[sheet_name]
    except KeyError as e:
        raise ValueError(
            f'AZN Template {template_path} has no sheet named "{sheet_name}". '
            f"Is this the right template?"
        ) from e
    except Exception as e:
        raise ValueError(f"Could not read AZN Template {template_path}: {e}") from e

    # ── Header cells ──────────────────────────────────────────────────────────
    ws[_azn["cell_ctr_ref"]] = f"CTR-26-{job_ref} AZN"
    ws[_azn["cell_job_ref"]] = int(job_ref)

    if header:
        if header.get("client"):      ws[_azn["cell_client"]]     = header["client"]
        if header.get("sub_client"):  ws[_azn["cell_sub_client"]] = header["sub_client"]
        if header.get("location"):    ws[_azn["cell_location"]]   = header["location"]
        if header.get("date"):        ws[_azn["cell_date"]]       = _parse_date(header["date"])
        if header.get("contract_no"): ws[_azn["cell_contract_no"]] = header["contract_no"]
        if header.get("revision") not in (None, ""):
            ws[_azn["cell_revision"]] = header["revision"]
        if header.get("scope"):       ws[_azn["cell_scope"]]      = header["scope"]

    # ── Insert extra rows before writing so overflow doesn't clobber the
    #    second section or the summary block ──────────────────────────────────
    _COLS = 9   # columns A–I used by data rows

    support_capacity = _SUPPORT_DATA_END - _SUPPORT_DATA_START + 1
    support_extra    = max(0, len(support_rows) - support_capacity)
    if support_extra:
        _insert_rows_preserving_merges(ws, _SUPPORT_DATA_END + 1, support_extra)
        for i in range(support_extra):
            _copy_row_format(ws, _SUPPORT_DATA_END, _SUPPORT_DATA_END + 1 + i, _COLS)

    # Second section has shifted down by support_extra
    other_start = _OTHER_DATA_START + support_extra
    other_end   = _OTHER_DATA_END   + support_extra

    other_capacity = other_end - other_start + 1
    other_extra    = max(0, len(other_rows) - other_capacity)
    if other_extra:
        _insert_rows_preserving_merges(ws, other_end + 1, other_extra)
        for i in range(other_extra):
            _copy_row_format(ws, other_end, other_end + 1 + i, _COLS)
    other_eff_end = other_end + other_extra

    # Total row-shift to apply to summary cell addresses
    _row_shift = support_extra + other_extra

    # ── Section-2 header text ("Onshore Activities" / "Offshore Activities"),
    #    driven by the CTR request's project type rather than fixed template
    #    text — the template ships with "Offshore Activities" hardcoded at
    #    A21/H21, which is wrong for an Onshore-only project. Written at the
    #    section's shifted title row (2 rows above its data start), to both
    #    the merged A-column cell and its standalone H-column mirror. ────────
    other_title_row = other_start - 2
    label = activities_label((header or {}).get("project_type", "") if header else "")
    ws.cell(row=other_title_row, column=1).value = label
    ws.cell(row=other_title_row, column=8).value = label

    # ── Write section helper (captures ws via closure) ─────────────────────────
    def _write_section(rows, start, eff_end, total_gap) -> int:
        """Writes one section's rows and its total; returns the total row number."""
        clear_end = max(eff_end, start + len(rows) - 1) if rows else eff_end
        for row in range(start, clear_end + 1):
            for col in range(1, _COLS + 1):
                ws.cell(row=row, column=col).value = None

        for i, lr in enumerate(rows):
            r = start + i
            try:
                num_emp = float(lr.get("num_employees", 1) or 1)
            except (ValueError, TypeError):
                num_emp = 1.0
            try:
                qty = float(lr.get("quantity", 0) or 0)
            except (ValueError, TypeError):
                qty = 0.0
            try:
                rate = float(lr.get("rate_azn", 0) or 0)
            except (ValueError, TypeError):
                rate = 0.0

            ws.cell(row=r, column=1).value = strip_support_keyword(str(lr.get("comment", "")))
            ws.cell(row=r, column=2).value = num_emp
            ws.cell(row=r, column=3).value = strip_support_keyword(str(lr.get("description", "")))
            ws.cell(row=r, column=4).value = qty
            ws.cell(row=r, column=5).value = str(lr.get("uom", "Hours"))
            ws.cell(row=r, column=6).value = rate
            ws.cell(row=r, column=7).value = f"=B{r}*D{r}*F{r}"
            ws.cell(row=r, column=8).value = f"=G{r}"
            ws.cell(row=r, column=9).value = str(lr.get("nationality", "NAT"))

        # The total always sits total_gap rows after the section's full
        # capacity block (eff_end), not after however many rows were
        # actually written — otherwise a request with fewer rows than the
        # template's capacity writes the total into a blank filler row
        # instead of the fixed "Total Project Support"/"Total Offshore" row.
        # Summing the whole capacity range (not just the written rows) is
        # safe — the cleared rows above are blank and contribute 0.
        total_row = eff_end + total_gap + 1
        ws.cell(row=total_row, column=7).value = f"=SUM(G{start}:G{eff_end})"
        return total_row

    support_total_row = _write_section(
        support_rows, _SUPPORT_DATA_START, _SUPPORT_DATA_END + support_extra, _SUPPORT_TOTAL_GAP)
    other_total_row = _write_section(
        other_rows,   other_start,         other_eff_end,                   _OTHER_TOTAL_GAP)

    # ── Summary cells — addresses from config, shifted by any inserted rows ────
    _r, _c = _cell_row_col(_azn["cell_summary_onshore"])
    support_summary_addr = f"{_col_letter(_c)}{_r + _row_shift}"
    ws.cell(row=_r + _row_shift, column=_c).value = f"=G{support_total_row}"

    _r, _c = _cell_row_col(_azn["cell_summary_offshore"])
    other_summary_addr = f"{_col_letter(_c)}{_r + _row_shift}"
    ws.cell(row=_r + _row_shift, column=_c).value = f"=G{other_total_row}"

    _r, _c = _cell_row_col(_azn["cell_summary_combined"])
    ws.cell(row=_r + _row_shift, column=_c).value = f"={support_summary_addr}+{other_summary_addr}"

    try:
        wb.save(out_path)
    except PermissionError as e:
        raise ValueError(
            f"Cannot save {out_path} — it may be open in Excel in another "
            f"window. Close it there and try again."
        ) from e
    finally:
        wb.close()
    return out_path
